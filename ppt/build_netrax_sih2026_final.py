# -*- coding: utf-8 -*-
"""NETRAX_SIH_2026 — final 6-slide SIH26106 idea-submission deck.
Chrome matches the official SIH template (oval team badge, idea-title
header, SIH hexagon mark, blue underlined subheadings, blue footer banner
with page number). All content is native editable PowerPoint text/shapes —
no slide is rasterized. Content is grounded in the repo's own docs
(README.md, docs/FINAL_STATUS.md, docs/AGENT_ARCHITECTURE.md,
docs/RISK_SCORING.md, docs/DATA_SOURCES.md, docs/ml/TRAINING_REPORT.md).
"""
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- THEME ----
NAVY        = "1F3864"
NAVY_DARK   = "16264A"
BLUE        = "1F4E9C"
BLUE_DK     = "16357A"
BLUE_LIGHT  = "D6E4F7"
BLUE_PALE   = "EEF3FC"
PURPLE      = "6B4C9A"
PURPLE_LT   = "E9E1F4"
TEAL_DK     = "0F6B63"
BLACK       = "0D0D0D"
TEXT        = "1A1A1A"
MUTED       = "44546A"
SOFT        = "6B7A99"
WHITE       = "FFFFFF"
BORDER      = "C9D3E3"
GREEN       = "1E8449"
GREEN_LT    = "E5F3EA"
AMBER       = "B7791F"
AMBER_LT    = "FBF0DC"
RED         = "B03A2E"
RED_LT      = "FBEAE8"
SERIF       = "Times New Roman"
SANS        = "Calibri"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

TEAM_NAME = "Nexora"
TEAM_ID = "[TEAM ID]"
PS_ID = "SIH26106"
PS_TITLE = "AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform"
THEME = "Blockchain & Cybersecurity"
CATEGORY = "Software"
ORG = "All India Council for Technical Education (Cyber Security Cell)"
DEPT = "Cyber Security Cell"

# --------------------------------------------------------------- HELPERS ---

def C(h):
    return RGBColor.from_string(h)


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = C(WHITE)
    bg.line.fill.background(); bg.shadow.inherit = False
    return slide


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.08,
             shape_type=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    no_shadow(shp)
    shp.text_frame.margin_left = 0; shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0; shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, w, h=None, fill=None, line=None, line_w=1.1):
    h = h if h is not None else w
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    no_shadow(shp)
    return shp


def add_line(slide, x1, y1, x2, y2, color=BORDER, weight=1.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = C(color)
    conn.line.width = Pt(weight)
    if dash:
        d = conn.line._get_or_add_ln()
        pd = etree.SubElement(d, qn('a:prstDash'))
        pd.set('val', dash)
    no_shadow(conn)
    return conn


def add_text(slide, x, y, w, h, text, size=12, color=TEXT, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             underline=False, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font; r.font.color.rgb = C(color); r.font.underline = underline
    return tb


def add_multi(slide, x, y, w, h, runs, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              line_spacing=1.05, font=SANS):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for text, color, bold, italic in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = font; r.font.color.rgb = C(color)
    return tb


def add_bullets(slide, x, y, w, h, items, size=12, color=TEXT, font=SANS, line_spacing=1.15,
                 space_after=8, bullet_color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r0 = p.add_run(); r0.text = "•  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = font
        r0.font.color.rgb = C(bullet_color or color)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.name = font
            r1.font.color.rgb = C(color)
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.name = font
            r2.font.color.rgb = C(color)
        else:
            r1 = p.add_run(); r1.text = item
            r1.font.size = Pt(size); r1.font.name = font
            r1.font.color.rgb = C(color)
    return tb


def sih_badge(slide, x, y):
    hexs = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(0.58), Inches(0.58))
    hexs.fill.solid(); hexs.fill.fore_color.rgb = C(NAVY)
    hexs.line.fill.background(); no_shadow(hexs)
    add_text(slide, x, y, 0.58, 0.58, "SIH", size=12, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x + 0.68, y - 0.03, 2.1, 0.32, "SMART INDIA\nHACKATHON", size=10, bold=True,
              color=NAVY, line_spacing=0.95)
    add_text(slide, x + 0.68, y + 0.29, 2.1, 0.2, "2026", size=10, bold=True, color=BLUE)


def team_badge(slide, name=TEAM_NAME):
    add_oval(slide, 0.35, 0.20, 1.15, 0.58, fill=WHITE, line=PURPLE, line_w=1.4)
    add_text(slide, 0.35, 0.20, 1.15, 0.58, name, size=11.5, color=TEXT, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)


def chrome(slide, title):
    team_badge(slide)
    sih_badge(slide, 10.95, 0.19)
    add_text(slide, 1.6, 0.18, 9.15, 0.6, title, size=23, bold=True, color=BLACK,
              align=PP_ALIGN.CENTER, font=SERIF, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
    add_rect(slide, 0.35, 0.90, SLIDE_W_IN - 0.7, 0.02, fill=BLUE)


def footer(slide, page, total=6):
    y = SLIDE_H_IN - 0.40
    add_rect(slide, 0, y, SLIDE_W_IN, 0.40, fill=BLUE_DK)
    add_text(slide, 0, y, SLIDE_W_IN - 0.5, 0.40, "NetraX  ·  SIH 2026 Idea Submission  ·  SIH26106",
              size=10.5, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, SLIDE_W_IN - 0.85, y, 0.6, 0.40, f"{page}/{total}", size=10.5, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def subhead(slide, x, y, w, text, size=16.5):
    add_text(slide, x, y, 0.26, 0.30, "❖", size=size, bold=True, color=BLUE)
    add_text(slide, x + 0.28, y, w - 0.28, 0.30, text, size=size, bold=True, color=BLUE,
              underline=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)


def circle_edge(cx, cy, r, dx, dy):
    n = math.hypot(dx, dy) or 1.0
    return cx + dx / n * r, cy + dy / n * r


def connect(slide, c1x, c1y, r1, c2x, c2y, r2, color=BORDER, weight=1.0, dash=None):
    dx, dy = c2x - c1x, c2y - c1y
    x1, y1 = circle_edge(c1x, c1y, r1, dx, dy)
    x2, y2 = circle_edge(c2x, c2y, r2, -dx, -dy)
    add_line(slide, x1, y1, x2, y2, color=color, weight=weight, dash=dash)


def node(slide, cx, cy, r, icon, fill, icon_color=WHITE, icon_size=None, line=None):
    add_oval(slide, cx - r, cy - r, r * 2, fill=fill, line=line, line_w=1.3)
    add_text(slide, cx - r, cy - r, r * 2, r * 2, icon, size=icon_size or int(r * 30),
              color=icon_color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def right_arrow(slide, x, y_center, size=0.16, color=SOFT):
    w = size; h = size * 0.75
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y_center - h / 2),
                                  Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = C(color)
    shp.line.fill.background(); no_shadow(shp)
    try:
        shp.adjustments[0] = 0.55; shp.adjustments[1] = 0.55
    except Exception:
        pass


# ============================================================== SLIDE 1 ====

def build_slide1(prs):
    slide = add_slide(prs)
    team_badge(slide)
    sih_badge(slide, 10.95, 0.19)

    add_text(slide, 0, 0.88, SLIDE_W_IN, 0.62, "NETRAX", size=40, bold=True, color=NAVY,
              align=PP_ALIGN.CENTER, font=SERIF)
    add_text(slide, 1.4, 1.52, SLIDE_W_IN - 2.8, 0.45,
              "AI-Powered Email Threat Detection, GeoLocation & Forensic Intelligence Platform",
              size=14.5, bold=True, color=MUTED, align=PP_ALIGN.CENTER, line_spacing=1.05)
    add_text(slide, 0, 1.98, SLIDE_W_IN, 0.30,
              "“Others detect the threat. NetraX investigates the threat.”",
              size=12.5, italic=True, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, font=SERIF)

    # hero investigation flow (editable shapes)
    fy = 2.55
    stages = [("✉", "Suspicious\nEmail"), ("\U0001F50D", "AI\nInvestigation"),
              ("\U0001F4C4", "Forensic\nEvidence"), ("\U0001F310", "Infrastructure\nIntelligence"),
              ("✓", "Risk /\nAction")]
    n = len(stages)
    r = 0.36
    margin = 1.9
    usable = SLIDE_W_IN - 2 * margin
    xs = [margin + usable * (i / (n - 1)) for i in range(n)]
    cy = fy + r
    for i in range(n - 1):
        add_line(slide, xs[i] + r, cy, xs[i + 1] - r, cy, color=BORDER, weight=1.2)
    for i, (icon, label) in enumerate(stages):
        col = BLUE if i not in (2, 3) else NAVY
        if i == 4:
            col = GREEN
        node(slide, xs[i], cy, r, icon, fill=col, icon_size=15)
        add_text(slide, xs[i] - 0.85, cy + r + 0.08, 1.7, 0.38, label, size=9.5, bold=True,
                  color=NAVY, align=PP_ALIGN.CENTER, line_spacing=0.95)

    # info block
    iy = 3.70
    add_rect(slide, 0.7, iy, SLIDE_W_IN - 1.4, 1.55, fill=BLUE_PALE, line=BORDER, line_w=1.0,
             radius=0.05, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    row1 = [("Problem Statement ID", PS_ID), ("Theme", THEME), ("PS Category", CATEGORY)]
    row2 = [("Team ID", TEAM_ID), ("Team Name", TEAM_NAME), ("Department", DEPT)]
    col_w = (SLIDE_W_IN - 1.4 - 0.6) / 3
    x0 = 0.7 + 0.3
    ry = iy + 0.16
    for ci, (lbl, val) in enumerate(row1):
        cx = x0 + ci * col_w
        add_text(slide, cx, ry, col_w - 0.1, 0.18, lbl.upper(), size=8.5, bold=True, color=SOFT)
        add_text(slide, cx, ry + 0.20, col_w - 0.1, 0.24, val, size=12, bold=True, color=NAVY)
    ry2 = ry + 0.52
    for ci, (lbl, val) in enumerate(row2):
        cx = x0 + ci * col_w
        add_text(slide, cx, ry2, col_w - 0.1, 0.18, lbl.upper(), size=8.5, bold=True, color=SOFT)
        add_text(slide, cx, ry2 + 0.20, col_w - 0.1, 0.24, val, size=12, bold=True, color=NAVY)
    add_line(slide, x0, ry2 + 0.50, SLIDE_W_IN - 0.7 - 0.3, ry2 + 0.50, color=BORDER, weight=0.75)
    add_text(slide, x0, ry2 + 0.56, SLIDE_W_IN - 1.4 - 0.6, 0.18, "PROBLEM STATEMENT TITLE",
              size=8.5, bold=True, color=SOFT)
    add_text(slide, x0, ry2 + 0.76, SLIDE_W_IN - 1.4 - 0.6, 0.22, PS_TITLE, size=11, bold=True, color=NAVY)

    add_text(slide, 0.7, 5.42, SLIDE_W_IN - 1.4, 0.24, ORG, size=11, bold=True, color=MUTED,
              align=PP_ALIGN.CENTER)
    add_text(slide, 0.7, 5.68, SLIDE_W_IN - 1.4, 0.4,
              "Decision-support & forensic-intelligence prototype — not a substitute for banks, "
              "police or official cybercrime reporting (cybercrime.gov.in).",
              size=9.5, italic=True, color=SOFT, align=PP_ALIGN.CENTER, line_spacing=1.15)
    footer(slide, 1)
    return slide


# ============================================================== SLIDE 2 ====

def build_slide2(prs):
    slide = add_slide(prs)
    chrome(slide, "NETRAX — PROPOSED SOLUTION & INNOVATION")
    left, cw = 0.55, SLIDE_W_IN - 1.1
    y = 1.12

    # 7-step flow
    chain = [("✉", "Suspicious\nEmail"), ("\U0001F4C4", "Email\nForensics"),
             ("\U0001F9E0", "Agentic\nInvestigation"), ("\U0001F6E1", "Threat\nIntelligence"),
             ("\U0001F310", "Infrastructure\nAnalysis"), ("\U0001F517", "Evidence\nCorrelation"),
             ("⚖", "Explainable\nRisk")]
    n = len(chain)
    r = 0.30
    margin = 0.85
    xs = [left + margin - left + (cw - 2 * (margin - left)) * (i / (n - 1)) for i in range(n)]
    cy = y + r
    for i in range(n - 1):
        add_line(slide, xs[i] + r, cy, xs[i + 1] - r, cy, color=BORDER, weight=1.1)
    for i, (icon, label) in enumerate(chain):
        col = BLUE if i != 2 else NAVY
        rr = r * 1.15 if i == 2 else r
        node(slide, xs[i], cy, rr, icon, fill=col, icon_size=13)
        add_text(slide, xs[i] - 0.62, cy + rr + 0.05, 1.24, 0.32, label, size=7.4, bold=True,
                  color=NAVY, align=PP_ALIGN.CENTER, line_spacing=0.9)

    y2 = cy + r + 0.48
    subhead(slide, left, y2, cw, "Proposed Solution")
    y2 += 0.36
    add_text(slide, left, y2, cw, 0.5,
              "NetraX is an agentic-AI platform that investigates a suspicious email beyond "
              "simple phishing detection — turning it into cited, correlated forensic evidence.",
              size=12.5, color=TEXT, line_spacing=1.15)
    y2 += 0.58

    subhead(slide, left, y2, cw, "How It Addresses the Problem")
    y2 += 0.36
    add_text(slide, left, y2, cw, 0.55,
              "It analyzes email content, headers, sender identity, URLs/domains, authentication "
              "signals, and source-IP infrastructure — then correlates that evidence to support "
              "forensic investigation, instead of returning a bare suspicious/not-suspicious label.",
              size=12.5, color=TEXT, line_spacing=1.15)
    y2 += 0.72

    subhead(slide, left, y2, cw, "Innovation & Uniqueness")
    y2 += 0.40
    col1 = [
        ("Agentic tool selection — ", "the agent dynamically chooses which investigation tools a case needs; nothing runs blindly."),
        ("Email forensics — ", "goes beyond message-content classification into headers, SPF/DKIM/DMARC and Received-chain analysis."),
        ("Threat intelligence — ", "adds external evidence via PhishTank and URLhaus correlation."),
    ]
    col2 = [
        ("Infrastructure context — ", "IP → ASN → approximate geolocation, restricted to public source IPs only."),
        ("Evidence Graph — ", "connects Email → Sender → Domain → URL → IP → ASN → Geo → Threat Intel."),
        ("Explainable risk — ", "a deterministic, per-source-capped score that supports — not replaces — investigator judgment."),
    ]
    colw = (cw - 0.4) / 2
    add_bullets(slide, left, y2, colw, 1.7, col1, size=11, line_spacing=1.15, space_after=8)
    add_bullets(slide, left + colw + 0.4, y2, colw, 1.7, col2, size=11, line_spacing=1.15, space_after=8)

    footer(slide, 2)
    return slide


# ============================================================== SLIDE 3 ====

def build_slide3(prs):
    slide = add_slide(prs)
    chrome(slide, "TECHNICAL APPROACH")
    left = 0.42
    content_w = SLIDE_W_IN - 0.84
    pipe_w = content_w * 0.60
    side_x = left + pipe_w + 0.28
    side_w = content_w - pipe_w - 0.28
    cx0 = left + pipe_w / 2
    top = 1.10

    def label(y, text, size=7.4):
        add_text(slide, left, y, pipe_w, 0.16, text, size=size, bold=True, color=SOFT)

    # ingestion
    chip_h = 0.28
    add_rect(slide, cx0 - 2.3, top, 4.6, chip_h, fill=NAVY, radius=0.5, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, cx0 - 2.3, top, 4.6, chip_h, "SUSPICIOUS EMAIL  →  INGESTION  →  PARSER",
              size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(top - 0.18, "01  INPUT")

    orch_r = 0.38
    orch_cy = top + chip_h + 0.28 + orch_r
    add_line(slide, cx0, top + chip_h, cx0, orch_cy - orch_r, color=BORDER, weight=1.0)
    add_oval(slide, cx0 - orch_r, orch_cy - orch_r, orch_r * 2, fill=BLUE)
    add_text(slide, cx0 - orch_r, orch_cy - orch_r + 0.04, orch_r * 2, 0.3, "NETRAX\nAGENT", size=7.6,
              bold=True, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=0.9)
    label(top + chip_h + 0.04, "02  AGENT — decides which tools this email needs")

    tools = [
        ("\U0001F4C4", "HEADERS\nFORENSICS", BLUE, False, ""),
        ("\U0001F517", "URL\nDOMAIN", BLUE, False, ""),
        ("\U0001F464", "SENDER\nIDENTITY", BLUE, False, ""),
        ("\U0001F6E1", "THREAT INTEL\n(if URL found)", TEAL_DK, True, "PhishTank · URLhaus"),
        ("\U0001F4CD", "GEOLOCATION\n(if public IP)", PURPLE, True, "Source IP → ASN → GeoLite2"),
    ]
    n_t = len(tools)
    t_r = 0.30
    t_gap = 0.36
    t_total = n_t * t_r * 2 + (n_t - 1) * t_gap
    tx0 = cx0 - t_total / 2 + t_r
    tool_y = orch_cy + orch_r + 0.60
    max_bottom = 0
    for i, (icon, lbl, col, cond, sub) in enumerate(tools):
        tx = tx0 + i * (t_r * 2 + t_gap)
        if cond:
            add_line(slide, cx0, orch_cy + orch_r, tx, tool_y - t_r, color=BORDER, weight=1.0, dash="dash")
            add_oval(slide, tx - t_r, tool_y - t_r, t_r * 2, fill=WHITE, line=col, line_w=1.5)
            add_text(slide, tx - t_r, tool_y - t_r, t_r * 2, t_r * 2, icon, size=10, color=col,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            connect(slide, cx0, orch_cy, orch_r, tx, tool_y, t_r, color=col, weight=1.2)
            node(slide, tx, tool_y, t_r, icon, fill=col, icon_size=10)
        add_text(slide, tx - 0.62, tool_y + t_r + 0.03, 1.24, 0.28, lbl, size=6.6, bold=True,
                  color=(SOFT if cond else NAVY), italic=cond, align=PP_ALIGN.CENTER, line_spacing=0.88)
        b = tool_y + t_r + 0.31
        if sub:
            add_text(slide, tx - 0.62, b, 1.24, 0.13, sub, size=5.6, italic=True, color=col,
                      align=PP_ALIGN.CENTER)
            b += 0.15
        max_bottom = max(max_bottom, b)
    label(tool_y - t_r - 0.18, "03  TOOLS")
    add_oval(slide, left, tool_y - 0.15, 0.09, fill=BLUE)
    add_text(slide, left + 0.13, tool_y - 0.185, 1.0, 0.14, "always run", size=5.8, italic=True, color=SOFT)
    add_oval(slide, left, tool_y + 0.06, 0.09, fill=WHITE, line=TEAL_DK, line_w=1.1)
    add_text(slide, left + 0.13, tool_y + 0.03, 1.0, 0.14, "conditional", size=5.8, italic=True, color=SOFT)

    evid_y = max_bottom + 0.18
    add_rect(slide, cx0 - 1.35, evid_y, 2.7, 0.24, fill=BLUE_PALE, line=BORDER, line_w=0.75,
             radius=0.5, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, cx0 - 1.35, evid_y, 2.7, 0.24, "EVIDENCE / INDICATORS", size=7.2, bold=True,
              color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(n_t):
        tx = tx0 + i * (t_r * 2 + t_gap)
        add_line(slide, tx, tool_y + t_r, cx0, evid_y, color=BORDER, weight=0.8)

    fus_r = 0.24
    fus_cy = evid_y + 0.24 + 0.20 + fus_r
    add_line(slide, cx0, evid_y + 0.24, cx0, fus_cy - fus_r, color=BORDER, weight=1.0)
    node(slide, cx0, fus_cy, fus_r, "\U0001F517", fill=NAVY_DARK, icon_size=9)
    add_text(slide, cx0 - 1.2, fus_cy + fus_r + 0.02, 2.4, 0.15, "EVIDENCE FUSION", size=7.2, bold=True,
              color=NAVY, align=PP_ALIGN.CENTER)

    risk_y = fus_cy + fus_r + 0.16 + 0.16
    add_line(slide, cx0, fus_cy + fus_r + 0.04, cx0, risk_y - 0.02, color=BORDER, weight=1.0)
    gauge_w, gauge_h = pipe_w - 1.7, 0.20
    gx = cx0 - gauge_w / 2
    seg_w = gauge_w / 4
    seg_cols = [GREEN, AMBER, "C0622A", RED]
    seg_lbl = ["LOW", "MEDIUM", "HIGH", "CRITICAL"]
    for i in range(4):
        shp_t = MSO_SHAPE.ROUNDED_RECTANGLE if i in (0, 3) else MSO_SHAPE.RECTANGLE
        add_rect(slide, gx + i * seg_w, risk_y, seg_w, gauge_h, fill=seg_cols[i], radius=0.5, shape_type=shp_t)
        add_text(slide, gx + i * seg_w, risk_y + gauge_h + 0.02, seg_w, 0.14, seg_lbl[i], size=6.0,
                  bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, left, risk_y - 0.02, gx - left - 0.08, gauge_h, "RISK ENGINE\n(deterministic) →",
              size=6.6, bold=True, color=NAVY, align=PP_ALIGN.RIGHT, line_spacing=0.9, anchor=MSO_ANCHOR.MIDDLE)

    dec_y = risk_y + gauge_h + 0.28
    decisions = [("\U0001F4AC", "EXPLAINABLE FINDINGS"), ("\U0001F5C2", "CASE REPORT")]
    d_w, d_gap = 2.35, 0.35
    dxp = cx0 - (d_w * 2 + d_gap) / 2
    for icon, lbl in decisions:
        add_rect(slide, dxp, dec_y, d_w, 0.30, fill=BLUE_PALE, line=BLUE_LIGHT, line_w=0.75, radius=0.5,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, dxp + 0.08, dec_y, 0.35, 0.30, icon, size=9, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, dxp + 0.40, dec_y, d_w - 0.45, 0.30, lbl, size=7.4, bold=True, color=BLUE_DK,
                  anchor=MSO_ANCHOR.MIDDLE)
        dxp += d_w + d_gap
    add_text(slide, left, dec_y, 1.6, 0.30, "04  DECIDE", size=6.8, bold=True, color=SOFT,
              anchor=MSO_ANCHOR.MIDDLE)

    # right sidebar
    sy = top
    subhead(slide, side_x, sy, side_w, "Datasets", size=12.5)
    sy += 0.32
    add_bullets(slide, side_x, sy, side_w, 0.95, [
        ("Enron Email Corpus — ", "language/structure reference only (not a malicious/benign label)."),
        ("SpamAssassin Corpus — ", "spam/ham email corpus; trains the email classifier."),
        ("UCI Phishing Websites — ", "trains the URL classifier."),
    ], size=8.6, line_spacing=1.08, space_after=4)
    sy += 1.02

    subhead(slide, side_x, sy, side_w, "Models — kept separate", size=12.5)
    sy += 0.32
    add_bullets(slide, side_x, sy, side_w, 0.85, [
        ("Email/text — ", "TF-IDF + Linear SVM. F1=0.961, ROC-AUC=0.998 (measured, held-out SpamAssassin set)."),
        ("URL/structured — ", "HistGradientBoosting. F1=0.954 (measured, UCI Phishing set)."),
    ], size=8.6, line_spacing=1.1, space_after=5)
    sy += 0.92
    add_text(slide, side_x, sy, side_w, 0.32,
              "Two independent, specialized models — not one combined model. The agentic layer "
              "orchestrates tool selection; the risk engine performs deterministic scoring.",
              size=8.2, italic=True, color=MUTED, line_spacing=1.15)
    sy += 0.55

    subhead(slide, side_x, sy, side_w, "Threat Intel & Geo", size=12.5)
    sy += 0.32
    add_bullets(slide, side_x, sy, side_w, 0.6, [
        ("PhishTank / URLhaus — ", "phishing / malware-URL threat intelligence."),
        ("MaxMind GeoLite2 — ", "IP → ASN → geolocation enrichment."),
    ], size=8.6, line_spacing=1.08, space_after=4)

    # bottom tech-stack strip
    strip_y = 6.40
    add_line(slide, left, strip_y - 0.08, SLIDE_W_IN - 0.42, strip_y - 0.08, color=BORDER, weight=0.75)
    stacks = [
        ("FRONTEND", "React · TypeScript · Vite · Tailwind CSS · shadcn/ui"),
        ("BACKEND", "Supabase · PostgreSQL · Edge Functions"),
        ("ML", "Python · Pandas · NumPy · scikit-learn"),
        ("AI / TI / GEO", "Gemini API (explanation, planned) · PhishTank · URLhaus · MaxMind GeoLite2"),
    ]
    sw = content_w / 4
    for i, (lbl, val) in enumerate(stacks):
        sx = left + i * sw
        add_text(slide, sx, strip_y, sw - 0.15, 0.15, lbl, size=6.6, bold=True, color=SOFT)
        add_text(slide, sx, strip_y + 0.16, sw - 0.15, 0.45, val, size=7.4, color=NAVY, line_spacing=1.05)

    footer(slide, 3)
    return slide


# ============================================================== SLIDE 4 ====

def build_slide4(prs):
    slide = add_slide(prs)
    chrome(slide, "FEASIBILITY AND VIABILITY")
    left, cw = 0.55, SLIDE_W_IN - 1.1
    y = 1.12

    subhead(slide, left, y, cw, "Why Is NetraX Feasible?")
    y += 0.38
    add_bullets(slide, left, y, cw, 0.85, [
        "Public research email and phishing datasets are available for model development and validation.",
        "Modular Python ML inference integrates with the investigation platform as a separate service.",
        "Threat intelligence and IP/ASN/geolocation are integrated through dedicated, testable adapters.",
        "Supabase/PostgreSQL stores cases, evidence, indicators and investigation events at scale.",
    ], size=11, line_spacing=1.1, space_after=4)
    y += 0.98

    add_text(slide, left, y, 6.0, 0.24, "Challenges", size=12.5, bold=True, underline=True, color=BLACK, font=SERIF)
    add_text(slide, left + 6.3, y, 5.9, 0.24, "Mitigation", size=12.5, bold=True, underline=True, color=BLACK, font=SERIF)
    y += 0.30
    pairs = [
        ("Data quality", "Preprocessing + validation before training (reports/data_quality_report.md)"),
        ("False positives / false negatives", "Model evaluation + decision-threshold tuning on held-out data"),
        ("Unknown / new attack patterns", "Multi-signal investigation (headers + URL + ML + threat intel) + feedback loop"),
        ("Threat intelligence / API availability", "Adapter architecture + cached/demo data; reports “unavailable,” never fabricated"),
        ("Security / malicious input", "Input validation + HTML sanitization + SSRF protection + secret management"),
        ("Scalability", "Modular services + asynchronous investigation steps"),
    ]
    for ch, mit in pairs:
        add_oval(slide, left, y + 0.06, 0.10, fill=RED)
        add_text(slide, left + 0.18, y, 5.9, 0.30, ch, size=10.5, bold=True, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
        add_oval(slide, left + 6.3, y + 0.06, 0.10, fill=GREEN)
        add_text(slide, left + 6.3 + 0.18, y, 5.7, 0.30, mit, size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.98)
        y += 0.335

    y += 0.16
    subhead(slide, left, y, cw, "Deployment Feasibility", size=14)
    y += 0.36
    steps = ["Email\nInput", "Analysis\nServices", "Evidence\nStore", "Investigation\nUI"]
    n = len(steps)
    bw, gap = 2.6, 0.5
    total = n * bw + (n - 1) * gap
    bx = left + (cw - total) / 2
    for i, s in enumerate(steps):
        add_rect(slide, bx, y, bw, 0.5, fill=BLUE_PALE, line=BLUE_LIGHT, line_w=1.0, radius=0.16,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, bx, y, bw, 0.5, s, size=10.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        if i < n - 1:
            right_arrow(slide, bx + bw + 0.06, y + 0.25, size=0.22, color=SOFT)
        bx += bw + gap
    footer(slide, 4)
    return slide


# ============================================================== SLIDE 5 ====

def build_slide5(prs):
    slide = add_slide(prs)
    chrome(slide, "IMPACT AND BENEFITS")
    left, cw = 0.55, SLIDE_W_IN - 1.1
    y = 1.10

    subhead(slide, left, y, cw, "Target Audience", size=13)
    y += 0.34
    audience = ["Government", "Banks & Financial Institutions", "Enterprises", "Educational Institutions",
                "SOC / Security Teams", "Digital Forensic Investigators"]
    gap = 0.16
    aw = (cw - gap * (len(audience) - 1)) / len(audience)
    ax = left
    for a in audience:
        add_rect(slide, ax, y, aw, 0.42, fill=PURPLE_LT, line=PURPLE, line_w=0.75, radius=0.5,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, ax + 0.04, y, aw - 0.08, 0.42, a, size=8.4, bold=True, color=PURPLE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.9)
        ax += aw + gap
    y += 0.62

    # central flow
    steps = ["NETRAX", "EMAIL\nINVESTIGATION", "CORRELATED\nEVIDENCE", "ACTIONABLE\nDECISION"]
    n = len(steps)
    bw, bgap = 2.7, 0.35
    total = n * bw + (n - 1) * bgap
    bx = left + (cw - total) / 2
    cols = [NAVY, BLUE, BLUE, GREEN]
    for i, s in enumerate(steps):
        add_rect(slide, bx, y, bw, 0.42, fill=cols[i], radius=0.5, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, bx, y, bw, 0.42, s, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.9)
        if i < n - 1:
            right_arrow(slide, bx + bw + 0.03, y + 0.21, size=0.20, color=SOFT)
        bx += bw + bgap
    y += 0.70

    cols3 = [
        ("Operational Benefits", BLUE, ["Faster suspicious-email triage", "Centralized investigation evidence",
                                          "Reduced manual correlation effort", "Explainable risk assessment"]),
        ("Forensic Benefits", NAVY, ["Header and sender analysis", "URL / domain relationships",
                                       "IP / ASN infrastructure context", "Evidence Graph for investigation"]),
        ("Security Benefits", TEAL_DK, ["Earlier identification of phishing indicators", "Threat-intelligence correlation",
                                          "Better incident investigation", "Structured case reporting"]),
    ]
    colw = (cw - 0.4 * 2) / 3
    for i, (title, col, items) in enumerate(cols3):
        cx = left + i * (colw + 0.4)
        add_text(slide, cx, y, colw, 0.24, title, size=11.5, bold=True, underline=True, color=col)
        add_bullets(slide, cx, y + 0.30, colw, 1.3, items, size=9.6, line_spacing=1.15, space_after=5, bullet_color=col)
    y += 1.62

    subhead(slide, left, y, cw, "Before vs. After NetraX", size=13)
    y += 0.36
    add_text(slide, left, y, 0.9, 0.26, "Before:", size=10.5, bold=True, color=RED)
    add_text(slide, left + 0.95, y, 11.0, 0.26,
              "Suspicious Email  →  Manual Checks  →  Scattered Evidence  →  Slow Investigation",
              size=10.5, color=TEXT)
    y += 0.30
    add_text(slide, left, y, 0.9, 0.26, "After:", size=10.5, bold=True, color=GREEN)
    add_text(slide, left + 0.95, y, 11.0, 0.26,
              "Suspicious Email  →  NetraX Agentic Investigation  →  Correlated Evidence  →  "
              "Forensic Risk  →  Case Report", size=10.5, color=TEXT)
    footer(slide, 5)
    return slide


# ============================================================== SLIDE 6 ====

def build_slide6(prs):
    slide = add_slide(prs)
    chrome(slide, "RESEARCH AND REFERENCES")
    left, cw = 0.55, SLIDE_W_IN - 1.1
    y = 1.15

    cats = [
        ("Datasets", BLUE, [("CMU Enron Email Dataset", "cs.cmu.edu/~enron"),
                              ("Apache SpamAssassin Public Corpus", "spamassassin.apache.org"),
                              ("UCI Phishing Websites Dataset", "archive.ics.uci.edu/dataset/327")]),
        ("Threat Intelligence", TEAL_DK, [("PhishTank", "phishtank.org"),
                                            ("URLhaus (abuse.ch)", "urlhaus.abuse.ch")]),
        ("Geolocation", PURPLE, [("MaxMind GeoLite2", "dev.maxmind.com/geoip/geolite2-free-geolocation-data")]),
    ]
    colw = (cw - 0.35 * 2) / 3
    ch = 1.9
    for i, (title, col, items) in enumerate(cats):
        cx = left + i * (colw + 0.35)
        add_rect(slide, cx, y, colw, ch, fill=WHITE, line=BORDER, line_w=1.0, radius=0.06,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_rect(slide, cx, y, colw, 0.05, fill=col, radius=0.5)
        add_text(slide, cx + 0.16, y + 0.14, colw - 0.32, 0.22, title, size=12.5, bold=True, color=NAVY)
        iy = y + 0.44
        for name, src in items:
            add_oval(slide, cx + 0.16, iy + 0.05, 0.08, fill=col)
            add_text(slide, cx + 0.32, iy, colw - 0.48, 0.20, name, size=9.6, bold=True, color=TEXT)
            add_text(slide, cx + 0.32, iy + 0.19, colw - 0.48, 0.16, src, size=8, italic=True, color=MUTED)
            iy += 0.42

    y += ch + 0.22
    subhead(slide, left, y, cw, "Research Areas", size=14)
    y += 0.38
    areas = ["Email phishing detection", "Email / header forensics", "URL phishing detection",
             "Threat intelligence correlation", "Agentic / tool-based investigation", "Explainable cybersecurity"]
    gap = 0.20
    aw = (cw - gap * 2) / 3
    for i, a in enumerate(areas):
        r_, c_ = divmod(i, 3)
        ax = left + c_ * (aw + gap)
        ay = y + r_ * 0.46
        add_rect(slide, ax, ay, aw, 0.36, fill=BLUE_PALE, line=BLUE_LIGHT, line_w=0.75, radius=0.5,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, ax + 0.08, ay, aw - 0.16, 0.36, a, size=9.6, bold=True, color=BLUE_DK,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
    y += 2 * 0.46 + 0.14
    add_text(slide, left, y, cw, 0.24,
              "Only verified, currently-used sources are listed above — no fabricated papers, DOIs or citations.",
              size=9.5, italic=True, color=MUTED)
    footer(slide, 6)
    return slide


def main():
    prs = new_presentation()
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    build_slide4(prs)
    build_slide5(prs)
    build_slide6(prs)
    prs.save("NETRAX_SIH_2026.pptx")
    print("Saved NETRAX_SIH_2026.pptx")


if __name__ == "__main__":
    main()
