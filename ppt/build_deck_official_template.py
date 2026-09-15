# -*- coding: utf-8 -*-
"""NetraX - SIH26106 deck, built to match the OFFICIAL SIH "Idea Submission"
PPT template chrome the user provided (oval team-name badge top-left, black
serif title, SIH badge top-right, blue underlined subheading, bullet body,
blue footer banner with page number). Content is the same repo-grounded
facts used in build_deck_netrax_email.py.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = "1F3864"
BLUE = "1F4E9C"
BLUE_BANNER = "1544A0"
PURPLE = "6B4C9A"
BLACK = "0D0D0D"
TEXT = "1A1A1A"
MUTED = "44546A"
WHITE = "FFFFFF"
GREEN = "1E8449"
ORANGE = "C0622A"
SERIF = "Times New Roman"
SANS = "Calibri"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

TEAM_NAME = "Nexora"
TEAM_ID = "[TEAM ID]"
PS_ID = "SIH26106"
PS_TITLE = "AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform"
THEME = "Blockchain & Cybersecurity"
CATEGORY = "Software"
ORG = "All India Council for Technical Education (Cyber Security Cell)"


def C(h):
    return RGBColor.from_string(h)


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def add_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = C(WHITE)
    bg.line.fill.background(); bg.shadow.inherit = False
    return slide


def no_line(shape):
    shape.line.fill.background()


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.0,
             shape_type=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    else:
        no_line(shp)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    no_shadow(shp)
    shp.text_frame.margin_left = 0; shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0; shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, w, h, fill=None, line=None, line_w=1.25):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    else:
        no_line(shp)
    no_shadow(shp)
    return shp


def add_text(slide, x, y, w, h, text, size=12, color=TEXT, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             underline=False, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font; r.font.color.rgb = C(color); r.font.underline = underline
    return tb


def add_bullets(slide, x, y, w, h, items, size=14, color=TEXT, font=SANS,
                 line_spacing=1.15, space_after=10, bullet_color=None, bold_lead=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r0 = p.add_run()
        r0.text = "•  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = font
        r0.font.color.rgb = C(bullet_color or color)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.name = font
            r1.font.color.rgb = C(color)
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.name = font
            r2.font.color.rgb = C(color)
        else:
            r1 = p.add_run(); r1.text = item
            r1.font.size = Pt(size); r1.font.name = font
            r1.font.color.rgb = C(color)
    return tb


def sih_badge(slide, x, y):
    """Simple generic SIH-2026 hexagon badge (not a reproduction of the official artwork)."""
    hexs = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
    hexs.fill.solid(); hexs.fill.fore_color.rgb = C(NAVY)
    no_line(hexs); no_shadow(hexs)
    add_text(slide, x, y, 0.62, 0.62, "SIH", size=13, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x + 0.72, y - 0.03, 2.1, 0.34, "SMART INDIA\nHACKATHON", size=10.5, bold=True,
              color=NAVY, line_spacing=0.95)
    add_text(slide, x + 0.72, y + 0.32, 2.1, 0.22, "2026", size=10.5, bold=True, color=BLUE)


def team_badge(slide, name=TEAM_NAME):
    ov = add_oval(slide, 0.35, 0.22, 1.15, 0.62, fill=WHITE, line=PURPLE, line_w=1.5)
    add_text(slide, 0.35, 0.22, 1.15, 0.62, name, size=12, bold=False, color=TEXT,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=SANS)


def chrome(slide, num_on_slide, title, subtitle=None):
    team_badge(slide)
    sih_badge(slide, 10.9, 0.20)
    add_text(slide, 1.7, 0.22, 9.0, 0.55, title, size=27, bold=True, color=BLACK,
              align=PP_ALIGN.CENTER, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 0.35, 0.98, SLIDE_W_IN - 0.7, 0.018, fill=BLUE)


def footer(slide, page, total=6):
    y = SLIDE_H_IN - 0.42
    add_rect(slide, 0, y, SLIDE_W_IN, 0.42, fill=BLUE_BANNER)
    add_text(slide, 0, y, SLIDE_W_IN - 0.5, 0.42, "NetraX  ·  SIH 2026 Idea Submission — SIH26106",
              size=11, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=SANS)
    add_text(slide, SLIDE_W_IN - 0.9, y, 0.7, 0.42, str(page), size=11, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def subheading(slide, x, y, w, text):
    add_text(slide, x, y, 0.30, 0.34, "❖", size=17, bold=True, color=BLUE, font=SANS)
    add_text(slide, x + 0.32, y, w - 0.32, 0.34, text, size=19, bold=True, color=BLUE,
              underline=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================== SLIDE 1 ====

def build_slide1(prs):
    slide = add_slide(prs)
    sih_badge(slide, 10.9, 0.28)

    labels = [
        ("Problem Statement ID", PS_ID),
        ("Problem Statement Title", PS_TITLE),
        ("Theme", THEME),
        ("PS Category", CATEGORY),
        ("Team ID", TEAM_ID),
        ("Team Name", TEAM_NAME),
    ]
    ly = 0.45
    for lbl, val in labels:
        add_text(slide, 0.55, ly, 2.9, 0.24, lbl, size=12, bold=True, underline=True, color=BLACK, font=SANS)
        add_text(slide, 0.55, ly + 0.24, 7.4, 0.4, val, size=12.5, color=TEXT, font=SANS, line_spacing=1.0)
        ly += 0.24 + (0.40 if len(val) > 70 else 0.30)

    add_oval(slide, 4.5, 2.55, 4.3, 4.3, fill="F4F1FA", line=PURPLE, line_w=1.5)
    add_text(slide, 4.5, 3.55, 4.3, 0.6, "NETRAX", size=44, bold=True, color=NAVY,
              align=PP_ALIGN.CENTER, font=SERIF)
    add_text(slide, 4.7, 4.20, 3.9, 0.7,
              "AI-Powered Email Threat Detection,\nGeoLocation & Forensic Intelligence",
              size=13, bold=True, color=MUTED, align=PP_ALIGN.CENTER, font=SANS, line_spacing=1.1)
    add_text(slide, 4.6, 5.55, 4.1, 0.4, "“Detection tells you WHAT.\nNetraX investigates WHY.”",
              size=11, italic=True, color=PURPLE, align=PP_ALIGN.CENTER, font=SERIF, line_spacing=1.1)

    add_text(slide, 0.55, 7.02, 12.2, 0.3, ORG, size=10.5, italic=True, color=MUTED, font=SANS)
    footer(slide, 1)
    return slide


# ============================================================== SLIDE 2 ====

def build_slide2(prs):
    slide = add_slide(prs)
    chrome(slide, 2, "PROPOSED SOLUTION", "")
    y = 1.20
    subheading(slide, 0.55, y, 12.2, "Proposed Solution (Idea / Prototype)")
    y += 0.48
    add_bullets(slide, 0.55, y, 12.2, 1.7, [
        ("Beyond detection: ", "NetraX investigates a suspicious email end-to-end — header forensics, URL/domain analysis, threat-intelligence correlation, and IP/ASN geolocation — before producing one explainable, deterministic risk score."),
        ("Addresses the problem: ", "traditional filters stop at “suspicious / not suspicious.” SOC teams and investigators are then left to manually piece together evidence. NetraX automates that investigation and hands back cited evidence, not just a verdict."),
    ], size=14, line_spacing=1.15, space_after=10)
    y += 1.55
    subheading(slide, 0.55, y, 12.2, "Innovation & Uniqueness")
    y += 0.46
    add_bullets(slide, 0.55, y, 12.2, 2.9, [
        ("Agentic investigation — ", "the orchestrator dynamically selects only the tools an email's own indicators call for (no URL → URL/threat-intel skipped; no public IP → geolocation skipped), with every skip logged as an auditable event, not a fixed always-run pipeline."),
        ("Email forensics — ", "sender identity, SPF/DKIM/DMARC, Received-chain and header anomalies, every finding cited to its literal evidence."),
        ("Threat intelligence — ", "extracted URLs/domains correlated against PhishTank and URLhaus where applicable."),
        ("Infrastructure intelligence — ", "public source IP → ASN → organization/network → approximate geolocation."),
        ("Evidence graph — ", "connects Email → Sender → Domain → URL → IP → ASN → Geo → Threat Intel."),
        ("Explainable risk — ", "multiple evidence signals fused into one transparent, deterministic risk score."),
    ], size=12.5, line_spacing=1.12, space_after=6)
    footer(slide, 2)
    return slide


# ============================================================== SLIDE 3 ====

def build_slide3(prs):
    slide = add_slide(prs)
    chrome(slide, 3, "TECHNICAL APPROACH", "")
    y = 1.20
    subheading(slide, 0.55, y, 12.2, "Investigation Pipeline")
    y += 0.46

    steps = ["Suspicious\nEmail", "Email\nParser", "NetraX Agent\n(orchestrator)", "Header / URL /\nSender Tools",
              "Threat Intel +\nGeolocation*", "Evidence\nFusion", "Risk Engine\n(0-100)", "Case\nReport"]
    n = len(steps)
    gap = 0.14
    bw = (12.2 - gap * (n - 1)) / n
    bx = 0.55
    for i, s in enumerate(steps):
        add_rect(slide, bx, y, bw, 0.62, fill="EAF0FB", line=BLUE, line_w=1.0, radius=0.12,
                 shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, bx + 0.03, y, bw - 0.06, 0.62, s, size=8.6, bold=True, color=NAVY,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        if i < n - 1:
            add_text(slide, bx + bw, y, gap, 0.62, "→", size=12, bold=True, color=BLUE,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        bx += bw + gap
    add_text(slide, 0.55, y + 0.66, 12.2, 0.22,
              "* run only when the email actually contains a URL / a public source IP — dynamic tool selection, not a fixed pipeline.",
              size=9.5, italic=True, color=MUTED)

    y += 1.05
    subheading(slide, 0.55, y, 6.0, "Data Grounding")
    y2 = y + 0.42
    add_bullets(slide, 0.55, y2, 6.0, 1.7, [
        ("Enron Email Corpus — ", "structural/language reference (not a labelled class)."),
        ("SpamAssassin Corpus — ", "trains the email content classifier."),
        ("UCI Phishing Websites — ", "trains the URL classifier."),
        ("PhishTank / URLhaus — ", "threat-intelligence adapters."),
        ("MaxMind GeoLite2 — ", "IP → ASN → geolocation."),
    ], size=11, line_spacing=1.1, space_after=4)

    subheading(slide, 6.85, y, 5.9, "Trained ML Models (measured)")
    y2 = y + 0.42
    add_bullets(slide, 6.85, y2, 5.9, 1.1, [
        ("Email classifier — ", "TF-IDF + Linear SVM, F1 = 0.961, ROC-AUC = 0.998 (held-out SpamAssassin test set)."),
        ("URL classifier — ", "HistGradientBoosting, F1 = 0.954 (UCI Phishing Websites test set)."),
    ], size=11, line_spacing=1.15, space_after=6)
    add_text(slide, 6.85, y2 + 1.05, 5.9, 0.5,
              "Risk score is computed by a deterministic, per-source-capped formula — an LLM never sets or adjusts it; it only orchestrates tool selection.",
              size=10.5, italic=True, color=MUTED, line_spacing=1.15)

    y3 = y2 + 1.75
    subheading(slide, 0.55, y3, 12.2, "Tech Stack")
    y3 += 0.42
    add_bullets(slide, 0.55, y3, 12.2, 0.9, [
        ("Frontend — ", "React 19, TypeScript, Vite, Tailwind CSS, shadcn/ui-style components, Framer Motion, Recharts."),
        ("Backend / ML — ", "Supabase / PostgreSQL, Edge Functions, Python + scikit-learn; Gemini API reserved for an explanation layer (not yet wired into scoring)."),
    ], size=11, line_spacing=1.1, space_after=4)
    footer(slide, 3)
    return slide


# ============================================================== SLIDE 4 ====

def build_slide4(prs):
    slide = add_slide(prs)
    chrome(slide, 4, "FEASIBILITY AND VIABILITY", "")
    y = 1.20
    subheading(slide, 0.55, y, 12.2, "Is This Idea Feasible?")
    y += 0.44
    add_text(slide, 0.55, y, 12.2, 0.55,
              "Yes — the core pipeline already runs end-to-end today: 203 automated tests passing "
              "(169 TypeScript + 34 Python, measured on this build) across parsing, forensics, threat "
              "intel, geolocation, the ML models, the agent orchestrator and the risk engine.",
              size=12.5, color=TEXT, line_spacing=1.2)
    y += 0.75

    add_text(slide, 0.55, y, 6.0, 0.24, "Challenges & Risks", size=13, bold=True, underline=True, color=BLACK, font=SANS)
    add_text(slide, 6.85, y, 5.9, 0.24, "Strategies for Overcoming", size=13, bold=True, underline=True, color=BLACK, font=SANS)
    y += 0.32
    pairs = [
        ("Data quality", "Preprocessing + profiling before training (reports/data_quality_report.md)"),
        ("Threat-intel / API availability", "Adapter architecture; reports “unavailable,” never a fabricated match — cached URLhaus snapshot for demo"),
        ("ML false positives / negatives", "Cross-validated model selection, tuned decision threshold, held-out evaluation"),
        ("Unseen attack patterns", "Multi-signal fusion (headers + URL + ML + threat intel) + a documented feedback loop"),
        ("Computational requirements", "ML inference isolated in its own service; investigation logic is portable/modular"),
        ("Platform security", "Auth + row-level security, input validation, SSRF-safe URL fetches, no secrets in code"),
    ]
    for ch, mit in pairs:
        add_text(slide, 0.55, y, 6.0, 0.34, "•  " + ch, size=11.5, bold=True, color=TEXT)
        add_text(slide, 6.85, y, 5.9, 0.34, "•  " + mit, size=11, color=MUTED, line_spacing=1.0)
        y += 0.36

    y += 0.10
    subheading(slide, 0.55, y, 12.2, "Why It Is Feasible")
    y += 0.42
    add_bullets(slide, 0.55, y, 12.2, 1.0, [
        "Public research datasets already acquired & profiled (Enron, SpamAssassin, UCI Phishing, URLhaus).",
        "PhishTank / URLhaus / GeoLite2 integrated via dedicated, unit-tested adapters.",
        "Supabase/PostgreSQL schema (14 tables, RLS) designed to scale case & evidence storage.",
        "Operates with live intelligence when credentialed, and clearly labelled demo data otherwise — never a fabricated result.",
    ], size=11.5, line_spacing=1.1, space_after=4)
    footer(slide, 4)
    return slide


# ============================================================== SLIDE 5 ====

def build_slide5(prs):
    slide = add_slide(prs)
    chrome(slide, 5, "IMPACT AND BENEFITS", "")
    y = 1.20
    add_text(slide, 0.55, y, 12.2, 0.5,
              "Target audience: Government, Banks / Financial Institutions, Enterprises, Educational "
              "Institutions, SOC / Security Teams, Digital Forensic Investigators.",
              size=12.5, bold=True, color=TEXT, line_spacing=1.15)
    y += 0.62

    cols = [
        ("SOC / Security Teams", GREEN, ["Faster triage & evidence correlation", "Tool-execution trace, fully auditable"]),
        ("Investigators", BLUE, ["Structured, citable forensic evidence", "Evidence graph across email/sender/URL/IP"]),
        ("Organizations", ORANGE, ["Better understanding of suspicious-email threats", "Reduced manual screening effort"]),
    ]
    cw = (12.2 - 0.4 * 2) / 3
    for i, (title, col, items) in enumerate(cols):
        cx = 0.55 + i * (cw + 0.4)
        add_text(slide, cx, y, cw, 0.26, title, size=13, bold=True, color=col, underline=True)
        add_bullets(slide, cx, y + 0.32, cw, 1.0, items, size=11, line_spacing=1.12, space_after=5, bullet_color=col)
    y += 1.55

    subheading(slide, 0.55, y, 12.2, "Before vs. After NetraX")
    y += 0.44
    before = "Suspicious Email  →  Manual Investigation  →  Scattered Evidence  →  Slow Decision"
    after = "Suspicious Email  →  NetraX Investigation  →  Correlated Evidence  →  Forensic Risk  →  Actionable Case Report"
    add_text(slide, 0.55, y, 1.0, 0.3, "Before:", size=12, bold=True, color=ORANGE)
    add_text(slide, 1.6, y, 11.0, 0.3, before, size=12, color=TEXT)
    y += 0.36
    add_text(slide, 0.55, y, 1.0, 0.3, "After:", size=12, bold=True, color=GREEN)
    add_text(slide, 1.6, y, 11.0, 0.3, after, size=12, color=TEXT)
    y += 0.55

    add_text(slide, 0.55, y, 12.2, 0.4,
              "A decision-support & forensic-intelligence platform — not a replacement for banks, "
              "police or official cybercrime systems.", size=11, italic=True, color=MUTED)
    footer(slide, 5)
    return slide


# ============================================================== SLIDE 6 ====

def build_slide6(prs):
    slide = add_slide(prs)
    chrome(slide, 6, "RESEARCH AND REFERENCES", "")
    y = 1.20

    subheading(slide, 0.55, y, 12.2, "Data & Intelligence")
    y += 0.42
    add_bullets(slide, 0.55, y, 12.2, 1.3, [
        ("CMU Enron Email Dataset — ", "cs.cmu.edu/~enron"),
        ("Apache SpamAssassin Public Corpus — ", "spamassassin.apache.org"),
        ("UCI Phishing Websites Dataset — ", "archive.ics.uci.edu/dataset/327"),
        ("PhishTank — ", "phishtank.org"),
        ("URLhaus (abuse.ch) — ", "urlhaus.abuse.ch"),
        ("MaxMind GeoLite2 — ", "dev.maxmind.com/geoip/geolite2-free-geolocation-data"),
    ], size=12, line_spacing=1.1, space_after=4)
    y += 1.55

    subheading(slide, 0.55, y, 6.0, "Technology")
    y2 = y + 0.42
    add_bullets(slide, 0.55, y2, 6.0, 1.6, [
        "React + TypeScript + Vite",
        "Tailwind CSS + shadcn/ui",
        "Supabase / PostgreSQL",
        "Python + scikit-learn",
        "Gemini API (explanation layer — planned)",
    ], size=11.5, line_spacing=1.1, space_after=4)

    subheading(slide, 6.85, y, 5.9, "Security & Standards")
    add_bullets(slide, 6.85, y2, 5.9, 1.6, [
        "OWASP Top 10 / SSRF Prevention — owasp.org",
        "NIST Cybersecurity Framework — nist.gov/cyberframework",
        "CERT-In Advisories — cert-in.org.in",
        "Govt. of India Cybercrime Portal — cybercrime.gov.in",
    ], size=11.5, line_spacing=1.1, space_after=4)

    y3 = y2 + 1.75
    subheading(slide, 0.55, y3, 12.2, "Research Basis")
    y3 += 0.42
    add_text(slide, 0.55, y3, 12.2, 0.6,
              "Email/Phishing Detection · URL Phishing Detection · Email Header Forensics "
              "(SPF/DKIM/DMARC) · Threat-Intelligence Correlation · Agentic / Tool-Use AI Systems · "
              "IP Geolocation Limitations · Explainable Security Systems",
              size=11.5, italic=True, color=MUTED, line_spacing=1.2)
    y3 += 0.7
    add_text(slide, 0.55, y3, 12.2, 0.25,
              "Topic areas grounding this build — specific citations to be finalized for the final report.",
              size=10, italic=True, color=MUTED)
    footer(slide, 6)
    return slide


def main():
    prs = new_presentation()
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    build_slide4(prs)
    build_slide5(prs)
    build_slide6(prs)
    out = "NetraX_SIH26106_Official_Template.pptx"
    prs.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
