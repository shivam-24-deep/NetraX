# -*- coding: utf-8 -*-
"""NetraX - SIH 2026 deck, v2 (advanced visual storytelling)."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- THEME ----
NAVY_DARK   = "0A1B33"
NAVY        = "10233E"
NAVY_MED    = "1C3A63"
NAVY_SOFT   = "27466F"
BLUE        = "2563EB"
BLUE_DK     = "1D4ED8"
BLUE_LIGHT  = "DBEAFE"
BLUE_PALE   = "EFF6FF"
TEAL        = "0D9488"
TEAL_DK     = "0F766E"
TEAL_LIGHT  = "CCFBF1"
GREEN       = "16A34A"
GREEN_LIGHT = "DCFCE7"
AMBER       = "D97706"
AMBER_LIGHT = "FEF3C7"
RED         = "DC2626"
RED_DK      = "991B1B"
RED_LIGHT   = "FEE2E2"
BG          = "F5F8FC"
CARD        = "FFFFFF"
BORDER      = "E2E8F0"
BORDER_DK   = "CBD5E1"
TEXT_DARK   = "0F172A"
TEXT_MUTED  = "55647E"
TEXT_SOFT   = "8592A8"
WHITE       = "FFFFFF"
FONT        = "Segoe UI"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

PLACEHOLDERS = {
    "ps_id": "[PROBLEM STATEMENT ID]",
    "ps_title": "[PROBLEM STATEMENT TITLE]",
    "theme": "[THEME]",
    "ps_category": "Software",
    "team_id": "[TEAM ID]",
    "team_name": "[TEAM NAME]",
}

# --------------------------------------------------------------- HELPERS ---

def C(hex_color):
    return RGBColor.from_string(hex_color)


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def add_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_fill(bg, BG)
    no_line(bg)
    bg.shadow.inherit = False
    return slide


def set_fill(shape, hex_color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(hex_color)


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, hex_color, width=0.75):
    shape.line.color.rgb = C(hex_color)
    shape.line.width = Pt(width)


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_soft_shadow(shape, blur_pt=9, dist_pt=2.2, dir_deg=90, alpha_pct=20, color="1E293B"):
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    effectLst = etree.SubElement(spPr, qn('a:effectLst'))
    outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
    outerShdw.set('blurRad', str(int(blur_pt * 12700)))
    outerShdw.set('dist', str(int(dist_pt * 12700)))
    outerShdw.set('dir', str(int(dir_deg * 60000)))
    outerShdw.set('rotWithShape', '0')
    clr = etree.SubElement(outerShdw, qn('a:srgbClr'))
    clr.set('val', color)
    alpha = etree.SubElement(clr, qn('a:alpha'))
    alpha.set('val', str(int(alpha_pct * 1000)))


def add_rect(slide, x, y, w, h, fill=CARD, line=None, line_w=0.75, radius=0.07,
             shadow=False, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill)
    if line:
        set_line(shp, line, line_w)
    else:
        no_line(shp)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    no_shadow(shp)
    if shadow:
        add_soft_shadow(shp)
    shp.text_frame.margin_left = 0
    shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0
    shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, d, fill=BLUE, line=None, line_w=1.0, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    if fill is None:
        shp.fill.background()
    else:
        set_fill(shp, fill)
    if line:
        set_line(shp, line, line_w)
    else:
        no_line(shp)
    no_shadow(shp)
    if shadow:
        add_soft_shadow(shp)
    return shp


def add_line_shape(slide, x1, y1, x2, y2, color=BORDER, weight=1.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = C(color)
    conn.line.width = Pt(weight)
    if dash:
        d = conn.line._get_or_add_ln()
        pd = etree.SubElement(d, qn('a:prstDash'))
        pd.set('val', dash)
    no_shadow(conn)
    return conn


def add_text(slide, x, y, w, h, text, size=12, color=TEXT_DARK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             wrap=True, spacing_pt=0.0, caps=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text.upper() if caps else text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = C(color)
    if spacing_pt:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(int(spacing_pt * 100)))
    return tb


def add_bullets(slide, x, y, w, h, items, size=10.5, color=TEXT_DARK, bullet_color=BLUE,
                 line_spacing=1.1, space_after=5, font=FONT, anchor=MSO_ANCHOR.TOP,
                 bullet_char="•"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r0 = p.add_run()
        r0.text = bullet_char + "  "
        r0.font.size = Pt(size)
        r0.font.bold = True
        r0.font.name = font
        r0.font.color.rgb = C(bullet_color)
        r1 = p.add_run()
        r1.text = item
        r1.font.size = Pt(size)
        r1.font.name = font
        r1.font.color.rgb = C(color)
    return tb


def add_shield(slide, x, y, w, h, fill=BLUE, line=None, check=True, check_color=WHITE,
                check_size=None):
    pts = [(0.0, 0.10), (0.5, 0.0), (1.0, 0.10), (1.0, 0.50), (0.5, 1.0), (0.0, 0.50)]
    abs_pts = [(Inches(x + px * w), Inches(y + py * h)) for px, py in pts]
    fb = slide.shapes.build_freeform(abs_pts[0][0], abs_pts[0][1])
    fb.add_line_segments(abs_pts[1:], close=True)
    shp = fb.convert_to_shape()
    set_fill(shp, fill)
    if line:
        set_line(shp, line, 1.0)
    else:
        no_line(shp)
    no_shadow(shp)
    if check:
        cs = check_size or (h * 0.5)
        add_text(slide, x, y + h * 0.16, w, cs, "✓", size=int(cs * 62), color=check_color,
                  bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shp


def add_logo(slide, x, y, size=0.34, wordmark=True, on_dark=True, word_size=13):
    add_shield(slide, x, y, size, size, fill=BLUE if on_dark else NAVY, check=True,
               check_color=WHITE, check_size=size)
    if wordmark:
        tcol = WHITE if on_dark else NAVY_DARK
        add_text(slide, x + size + 0.11, y - 0.02, 3.2, size + 0.06, "NETRAX",
                  size=word_size, bold=True, color=tcol, anchor=MSO_ANCHOR.MIDDLE,
                  spacing_pt=0.3)


def header_band(slide, num, title, subtitle=None, height=0.86):
    add_rect(slide, 0, 0, SLIDE_W_IN, height, fill=NAVY_DARK, shape_type=MSO_SHAPE.RECTANGLE)
    add_rect(slide, 0, height, SLIDE_W_IN, 0.04, fill=BLUE, shape_type=MSO_SHAPE.RECTANGLE)
    add_logo(slide, 0.42, height / 2 - 0.15, size=0.30, wordmark=True, on_dark=True, word_size=13.5)
    # slide number badge
    add_oval(slide, 3.55, height / 2 - 0.17, 0.34, fill=BLUE)
    add_text(slide, 3.55, height / 2 - 0.17, 0.34, 0.34, str(num), size=13, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ty = 0.13 if subtitle else height / 2 - 0.20
    add_text(slide, 4.05, ty, 6.6, 0.44, title, size=21, bold=True, color=WHITE, caps=True,
              spacing_pt=0.4, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 4.05, ty + 0.40, 6.6, 0.28, subtitle, size=11, color=BLUE_LIGHT, italic=True)
    add_text(slide, 10.55, 0.14, 2.35, 0.22, "SMART INDIA HACKATHON", size=9, bold=True,
              color=TEAL_LIGHT, align=PP_ALIGN.RIGHT, spacing_pt=0.4)
    add_text(slide, 10.55, 0.37, 2.35, 0.28, "2026  ·  SOFTWARE", size=12, bold=True,
              color=WHITE, align=PP_ALIGN.RIGHT)


def footer(slide, page_num, total=6, question=""):
    y = 7.20
    add_line_shape(slide, 0.42, y, SLIDE_W_IN - 0.42, y, color=BORDER_DK, weight=0.75)
    add_shield(slide, 0.42, y + 0.075, 0.15, 0.15, fill=NAVY, check=True, check_color=WHITE,
               check_size=0.15)
    add_text(slide, 0.64, y + 0.05, 5.5, 0.22, "NETRAX  ·  Detect · Investigate · Explain · Act",
              size=8, color=TEXT_MUTED, anchor=MSO_ANCHOR.MIDDLE)
    if question:
        add_text(slide, 4.5, y + 0.05, 5.3, 0.22, question, size=8, italic=True, color=TEXT_SOFT,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 9.5, y + 0.05, 3.38, 0.22, f"Page {page_num} / {total}",
              size=8, color=TEXT_MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def risk_dot(slide, x, y, d, color, line=None, line_w=1.0):
    shp = add_oval(slide, x, y, d, fill=color, line=line, line_w=line_w)
    return shp


def right_arrow(slide, x, y_center, size=0.16, color=BORDER_DK):
    w = size
    h = size * 0.78
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y_center - h / 2),
                                  Inches(w), Inches(h))
    set_fill(shp, color)
    no_line(shp)
    no_shadow(shp)
    try:
        shp.adjustments[0] = 0.55
        shp.adjustments[1] = 0.55
    except Exception:
        pass
    return shp


def down_arrow(slide, x_center, y, size=0.16, color=BORDER_DK):
    w = size * 1.3
    h = size
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x_center - w / 2), Inches(y),
                                  Inches(w), Inches(h))
    set_fill(shp, color)
    no_line(shp)
    no_shadow(shp)
    try:
        shp.adjustments[0] = 0.55
        shp.adjustments[1] = 0.55
    except Exception:
        pass
    return shp


def section_label(slide, x, y, text, color=NAVY_DARK, size=12.5, accent=BLUE, num=None):
    add_rect(slide, x, y + 0.03, 0.07, 0.20, fill=accent, radius=0.5)
    add_text(slide, x + 0.16, y, 7.0, 0.26, text, size=size, bold=True, color=color, caps=True,
              spacing_pt=0.3)


def icon_circle(slide, cx, cy, d, icon, fill, icon_size=None, icon_color=WHITE):
    add_oval(slide, cx - d / 2, cy - d / 2, d, fill=fill)
    add_text(slide, cx - d / 2, cy - d / 2, d, d, icon, size=icon_size or int(d * 32),
              color=icon_color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ---- geometry: connector edge-point helpers -------------------------------

def circle_edge_point(cx, cy, r, dx, dy):
    n = math.hypot(dx, dy) or 1.0
    ux, uy = dx / n, dy / n
    return cx + ux * r, cy + uy * r


def rect_edge_point(cx, cy, hw, hh, dx, dy):
    if dx == 0 and dy == 0:
        return cx, cy
    tx = hw / abs(dx) if dx != 0 else float('inf')
    ty = hh / abs(dy) if dy != 0 else float('inf')
    t = min(tx, ty)
    return cx + dx * t, cy + dy * t


def connect_circle_to_rect(slide, c_cx, c_cy, c_r, r_cx, r_cy, r_hw, r_hh, color=BORDER_DK,
                            weight=1.1, dash=None, inset=0.0):
    dx, dy = r_cx - c_cx, r_cy - c_cy
    x1, y1 = circle_edge_point(c_cx, c_cy, c_r + inset, dx, dy)
    x2, y2 = rect_edge_point(r_cx, r_cy, r_hw + inset, r_hh + inset, -dx, -dy)
    add_line_shape(slide, x1, y1, x2, y2, color=color, weight=weight, dash=dash)


def connect_circle_to_circle(slide, c1x, c1y, r1, c2x, c2y, r2, color=BORDER_DK, weight=1.1,
                              dash=None):
    dx, dy = c2x - c1x, c2y - c1y
    x1, y1 = circle_edge_point(c1x, c1y, r1, dx, dy)
    x2, y2 = circle_edge_point(c2x, c2y, r2, -dx, -dy)
    add_line_shape(slide, x1, y1, x2, y2, color=color, weight=weight, dash=dash)


def add_circuit_pattern(slide, x0, y0, cols, rows, spacing, dot_d=0.035, color=NAVY_SOFT):
    centers = []
    for r in range(rows):
        for c in range(cols):
            cx = x0 + c * spacing
            cy = y0 + r * spacing
            centers.append((cx, cy))
            risk_dot(slide, cx - dot_d / 2, cy - dot_d / 2, dot_d, color)
    for r in range(rows):
        for c in range(cols - 1):
            i = r * cols + c
            x1, y1 = centers[i]
            x2, y2 = centers[i + 1]
            add_line_shape(slide, x1, y1, x2, y2, color=color, weight=0.75)


# ============================================================== SLIDE 1 ====
# WHAT IS IT?  -- converging signal network into an AI core

def build_slide1(prs):
    slide = add_slide(prs)
    add_rect(slide, 0, 0, SLIDE_W_IN, SLIDE_H_IN, fill=NAVY_DARK, shape_type=MSO_SHAPE.RECTANGLE)
    add_rect(slide, 9.6, -1.6, 6.2, 6.2, fill=NAVY, shape_type=MSO_SHAPE.OVAL)
    add_rect(slide, -2.0, 6.2, 5.0, 5.0, fill=NAVY, shape_type=MSO_SHAPE.OVAL)
    add_circuit_pattern(slide, 11.75, 0.26, 3, 3, 0.24, dot_d=0.032, color=NAVY_SOFT)

    add_text(slide, 0, 0.26, SLIDE_W_IN, 0.28, "SMART INDIA HACKATHON 2026", size=12.5, bold=True,
              color=TEAL_LIGHT, align=PP_ALIGN.CENTER, spacing_pt=2.4)

    # ---- network diagram zone ----
    zone_top, zone_h = 0.72, 2.85
    cx0, cy0 = 6.35, zone_top + zone_h / 2
    core_r = 0.62

    inputs = [
        ("✉", "SMS"), ("@", "EMAIL"), ("\U0001F517", "URL"),
        ("☎", "PHONE"), ("\U0001F4B3", "TRANSACTION"),
    ]
    in_x = 1.45
    in_r = 0.26
    n_in = len(inputs)
    for i, (icon, label) in enumerate(inputs):
        iy = zone_top + (zone_h / n_in) * (i + 0.5)
        connect_circle_to_circle(slide, in_x, iy, in_r, cx0, cy0, core_r, color="2E4E7E",
                                   weight=1.1)
        add_oval(slide, in_x - in_r, iy - in_r, in_r * 2, fill=NAVY_SOFT, line="3B5D8A", line_w=1.0)
        add_text(slide, in_x - in_r, iy - in_r, in_r * 2, in_r * 2, icon, size=12, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, bold=True)
        add_text(slide, in_x + in_r + 0.12, iy - 0.11, 1.55, 0.22, label, size=8.6, bold=True,
                  color=BLUE_LIGHT, anchor=MSO_ANCHOR.MIDDLE, spacing_pt=0.5)

    outputs = [("RISK", AMBER), ("EVIDENCE", BLUE), ("ACTION", GREEN)]
    out_x = 11.15
    out_r = 0.34
    n_out = len(outputs)
    for i, (label, col) in enumerate(outputs):
        oy = zone_top + (zone_h / n_out) * (i + 0.5)
        connect_circle_to_circle(slide, cx0, cy0, core_r, out_x, oy, out_r, color="2E4E7E",
                                   weight=1.1)
        add_oval(slide, out_x - out_r, oy - out_r, out_r * 2, fill=col)
        add_text(slide, out_x - 0.9, oy + out_r + 0.03, 1.8, 0.18, label, size=7.8, bold=True,
                  color=WHITE, align=PP_ALIGN.CENTER, spacing_pt=0.5)

    # core
    add_oval(slide, cx0 - core_r - 0.16, cy0 - core_r - 0.16, (core_r + 0.16) * 2, fill=NAVY_MED)
    add_shield(slide, cx0 - core_r * 0.72, cy0 - core_r * 0.78, core_r * 1.44, core_r * 1.44,
               fill=BLUE, check=True, check_color=WHITE, check_size=core_r * 1.44)
    add_text(slide, cx0 - 1.1, cy0 + core_r + 0.24, 2.2, 0.22, "AI FRAUD CORE", size=9, bold=True,
              color=TEAL_LIGHT, align=PP_ALIGN.CENTER, spacing_pt=0.8)

    add_text(slide, in_x - 0.4, zone_top - 0.30, 1.8, 0.22, "SIGNALS IN", size=8, bold=True,
              color=TEXT_SOFT, spacing_pt=1.2)
    add_text(slide, out_x - 0.9, zone_top - 0.30, 1.8, 0.22, "OUTCOMES", size=8, bold=True,
              color=TEXT_SOFT, align=PP_ALIGN.CENTER, spacing_pt=1.2)

    # ---- title block ----
    add_text(slide, 0, 3.70, SLIDE_W_IN, 0.62, "NETRAX", size=40, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, spacing_pt=1.5)
    add_text(slide, 0, 4.34, SLIDE_W_IN, 0.36, "“Detect. Investigate. Explain. Act.”",
              size=16.5, bold=True, italic=True, color=TEAL_LIGHT, align=PP_ALIGN.CENTER)
    add_text(slide, 1.7, 4.75, SLIDE_W_IN - 3.4, 0.3,
              "Agentic AI for Cyber-Fraud Detection & Explainable Risk Assessment",
              size=11.5, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)

    # ---- info band ----
    info_y, info_h = 5.24, 1.02
    add_rect(slide, 0.8, info_y, SLIDE_W_IN - 1.6, info_h, fill=NAVY_MED, radius=0.09,
             line="2C4B78", line_w=0.75)
    cols = [
        ("PROBLEM STATEMENT ID", PLACEHOLDERS["ps_id"]),
        ("PROBLEM STATEMENT TITLE", PLACEHOLDERS["ps_title"]),
        ("THEME", PLACEHOLDERS["theme"]),
    ]
    cols2 = [
        ("PS CATEGORY", PLACEHOLDERS["ps_category"]),
        ("TEAM ID", PLACEHOLDERS["team_id"]),
        ("TEAM NAME", PLACEHOLDERS["team_name"]),
    ]
    col_w = (SLIDE_W_IN - 1.6 - 0.6) / 3
    col_x0 = 0.8 + 0.3
    for row_i, rowdata in enumerate([cols, cols2]):
        ry = info_y + 0.14 + row_i * 0.42
        for ci, (lbl, val) in enumerate(rowdata):
            cxp = col_x0 + ci * col_w
            add_text(slide, cxp, ry, col_w - 0.1, 0.17, lbl, size=7.6, bold=True, color=TEAL_LIGHT,
                      spacing_pt=0.8)
            add_text(slide, cxp, ry + 0.185, col_w - 0.1, 0.22, val, size=11, bold=True,
                      color=WHITE)

    add_text(slide, 0.8, 6.44, SLIDE_W_IN - 1.6, 0.5,
              "Decision-support & early-warning platform — not a replacement for banks, police "
              "or official cybercrime systems  ·  Prototype (SIH 2026)", size=8.5,
              italic=True, color=TEXT_SOFT, align=PP_ALIGN.CENTER, line_spacing=1.15)
    return slide


# ============================================================== SLIDE 2 ====
# WHY DIFFERENT?  -- confusion -> clarity transformation + USP pillars

def build_slide2(prs):
    slide = add_slide(prs)
    header_band(slide, 2, "Proposed Solution", "Why NetraX Is Different")
    top = 1.02

    panel_gap = 0.30
    panel_w = (SLIDE_W_IN - 0.84 - panel_gap) / 2
    xL = 0.42
    xR = xL + panel_w + panel_gap
    panel_h = 2.62

    # ---- LEFT: confusion panel ----
    add_rect(slide, xL, top, panel_w, panel_h, fill=RED_LIGHT, radius=0.06, line="F3B9B9",
             line_w=0.75)
    add_text(slide, xL + 0.24, top + 0.14, panel_w - 0.4, 0.22, "CURRENT EXPERIENCE", size=10.5,
              bold=True, color=RED_DK, spacing_pt=0.5)
    add_text(slide, xL + 0.24, top + 0.36, panel_w - 0.4, 0.2, "A suspicious message arrives — the user is on their own.",
              size=8.8, italic=True, color="8A3A3A")

    cxL, cyL = xL + panel_w / 2, top + panel_h / 2 + 0.28
    add_oval(slide, cxL - 0.34, cyL - 0.34, 0.68, fill=WHITE, line=RED, line_w=1.25)
    add_text(slide, cxL - 0.34, cyL - 0.34, 0.68, 0.68, "\U0001F4AC", size=22,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    qs = [
        ("Is it real?", -1.55, -0.66),
        ("Is this link safe?", 0.90, -0.66),
        ("Should I pay?", -1.55, 0.40),
        ("What should I do?", 0.90, 0.40),
    ]
    for txt, dx, dy in qs:
        qx, qy = cxL + dx, cyL + dy
        add_oval(slide, qx, qy, 0.34, fill=RED)
        add_text(slide, qx, qy, 0.34, 0.34, "?", size=13, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, qx - 0.48, qy + 0.36, 1.30, 0.28, txt, size=7.8, bold=True, italic=True,
                  color=RED_DK, align=PP_ALIGN.CENTER, line_spacing=1.0)

    # ---- RIGHT: clarity panel ----
    add_rect(slide, xR, top, panel_w, panel_h, fill=TEAL_LIGHT, radius=0.06, line="9FDCD3",
             line_w=0.75)
    add_text(slide, xR + 0.24, top + 0.14, panel_w - 0.4, 0.22, "NETRAX", size=10.5,
              bold=True, color=TEAL_DK, spacing_pt=0.5)
    add_text(slide, xR + 0.24, top + 0.36, panel_w - 0.4, 0.2,
              "One input → multi-signal investigation → explainable outcome.",
              size=8.8, italic=True, color=TEAL_DK)

    steps = ["SUSPICIOUS\nINPUT", "AI\nINVESTIGATION", "EVIDENCE", "RISK", "ACTION"]
    scols = [NAVY, TEAL, NAVY_SOFT, AMBER, GREEN]
    n_s = len(steps)
    s_gap = 0.12
    s_w = (panel_w - 0.48 - s_gap * (n_s - 1)) / n_s
    s_h = 0.58
    s_y = top + 0.70
    sx = xR + 0.24
    ai_cx = None
    for i, lbl in enumerate(steps):
        shp = add_rect(slide, sx, s_y, s_w, s_h, fill=scols[i], radius=0.20)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.92
        parts = lbl.split("\n")
        r = p.add_run(); r.text = parts[0]
        r.font.size = Pt(7.6); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = C(WHITE)
        if len(parts) > 1:
            p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run(); r2.text = parts[1]
            r2.font.size = Pt(7.6); r2.font.bold = True; r2.font.name = FONT
            r2.font.color.rgb = C(WHITE)
        if i == 1:
            ai_cx = sx + s_w / 2
        if i < n_s - 1:
            ax = sx + s_w + 0.01
            tri = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(s_y + s_h / 2 - 0.06),
                                          Inches(s_gap - 0.02), Inches(0.12))
            set_fill(tri, BORDER_DK)
            no_line(tri)
            no_shadow(tri)
        sx += s_w + s_gap

    # dynamic dispatch fan-out under "AI INVESTIGATION"
    tools = ["Msg", "URL", "Txn", "Pattern", "Intel"]
    fan_y = s_y + s_h + 0.30
    add_line_shape(slide, ai_cx, s_y + s_h + 0.02, ai_cx, fan_y - 0.10, color=TEAL_DK, weight=1.0,
                    dash="sysDot")
    fan_w_total = 2.9
    fan_x0 = ai_cx - fan_w_total / 2
    fdot = 0.26
    for i, tname in enumerate(tools):
        fx = fan_x0 + i * (fan_w_total / (len(tools) - 1)) - fdot / 2
        add_line_shape(slide, ai_cx, fan_y - 0.10, fx + fdot / 2, fan_y, color=TEAL_DK, weight=0.85)
        add_oval(slide, fx, fan_y, fdot, fill=WHITE, line=TEAL, line_w=1.1)
        add_text(slide, fx - 0.25, fan_y + fdot + 0.02, fdot + 0.5, 0.18, tname, size=6.8,
                  bold=True, color=TEAL_DK, align=PP_ALIGN.CENTER)
    add_text(slide, xR + 0.24, fan_y + fdot + 0.24, panel_w - 0.48, 0.2,
              "Agent dispatches only the tools this case needs", size=7.6, italic=True,
              color=TEAL_DK, align=PP_ALIGN.CENTER)

    # transformation badge straddling the two panels
    badge_w, badge_h = 1.55, 0.42
    bx = SLIDE_W_IN / 2 - badge_w / 2
    by = top + panel_h / 2 - badge_h / 2
    add_rect(slide, bx, by, badge_w, badge_h, fill=NAVY_DARK, radius=0.5, shadow=True)
    add_text(slide, bx, by, badge_w, badge_h, "CONFUSION → CLARITY", size=8.6, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing_pt=0.3)

    # ---- USP pillars ----
    pil_y = top + panel_h + 0.22
    section_label(slide, 0.42, pil_y, "What Makes The Agent Different", accent=BLUE)
    cards_y = pil_y + 0.34
    card_h = 1.62
    gap = 0.20
    card_w = (SLIDE_W_IN - 0.84 - gap * 3) / 4
    x0 = 0.42

    # 1. AGENTIC - ring badge style
    cx = x0
    add_rect(slide, cx, cards_y, card_w, card_h, fill=CARD, radius=0.09, line=BORDER, line_w=0.75)
    ringc = cx + card_w / 2
    add_oval(slide, ringc - 0.38, cards_y + 0.18, 0.76, fill=None, line=TEAL, line_w=1.6)
    add_oval(slide, ringc - 0.30, cards_y + 0.26, 0.60, fill=TEAL)
    add_text(slide, ringc - 0.30, cards_y + 0.26, 0.60, 0.60, "\U0001F9E0", size=18,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, cx + 0.12, cards_y + 1.02, card_w - 0.24, 0.26, "AGENTIC", size=11.5,
              bold=True, color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.5)
    add_text(slide, cx + 0.16, cards_y + 1.30, card_w - 0.32, 0.3, "Dynamic tool selection per case",
              size=8, color=TEXT_MUTED, align=PP_ALIGN.CENTER, line_spacing=1.0)

    # 2. EXPLAINABLE - horizontal chip style
    cx = x0 + (card_w + gap)
    add_rect(slide, cx, cards_y, card_w, card_h, fill=NAVY_DARK, radius=0.09)
    add_rect(slide, cx + 0.16, cards_y + 0.20, 0.5, 0.5, fill=BLUE, radius=0.22)
    add_text(slide, cx + 0.16, cards_y + 0.20, 0.5, 0.5, "\U0001F50D", size=16,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, cx + 0.78, cards_y + 0.22, card_w - 0.95, 0.5, "EXPLAINABLE", size=11,
              bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
    add_text(slide, cx + 0.16, cards_y + 0.88, card_w - 0.32, 0.6,
              "Every score traces back to visible evidence", size=8, color=BLUE_LIGHT,
              line_spacing=1.1)

    # 3. ADAPTIVE - mini process-dot style
    cx = x0 + 2 * (card_w + gap)
    add_rect(slide, cx, cards_y, card_w, card_h, fill=CARD, radius=0.09, line=BORDER, line_w=0.75)
    dot_y = cards_y + 0.30
    dxs = [cx + 0.30, cx + card_w / 2, cx + card_w - 0.30]
    for k in range(2):
        add_line_shape(slide, dxs[k], dot_y, dxs[k + 1], dot_y, color=AMBER, weight=1.4)
    for dxp in dxs:
        add_oval(slide, dxp - 0.09, dot_y - 0.09, 0.18, fill=AMBER)
    add_text(slide, cx + 0.12, cards_y + 0.60, card_w - 0.24, 0.26, "ADAPTIVE", size=11.5,
              bold=True, color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.5)
    add_text(slide, cx + 0.16, cards_y + 0.90, card_w - 0.32, 0.6,
              "Investigation path adapts to each case, not a fixed script", size=8,
              color=TEXT_MUTED, align=PP_ALIGN.CENTER, line_spacing=1.1)

    # 4. ACTIONABLE - checklist style
    cx = x0 + 3 * (card_w + gap)
    add_rect(slide, cx, cards_y, card_w, card_h, fill=CARD, radius=0.09, line=BORDER, line_w=0.75)
    add_text(slide, cx + 0.16, cards_y + 0.14, card_w - 0.32, 0.24, "ACTIONABLE", size=11.5,
              bold=True, color=NAVY_DARK, spacing_pt=0.5)
    acts = ["Do not click", "Do not share OTP", "Verify officially"]
    for i, a in enumerate(acts):
        ay = cards_y + 0.50 + i * 0.32
        add_oval(slide, cx + 0.16, ay, 0.20, fill=GREEN)
        add_text(slide, cx + 0.16, ay, 0.20, 0.20, "✓", size=9, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx + 0.44, ay - 0.015, card_w - 0.6, 0.24, a, size=8.4, color=TEXT_DARK,
                  anchor=MSO_ANCHOR.MIDDLE)

    # closing statement bar
    bar_y = cards_y + card_h + 0.16
    add_rect(slide, 0.42, bar_y, SLIDE_W_IN - 0.84, 0.46, fill=BLUE_PALE, radius=0.5, line=BLUE_LIGHT,
             line_w=0.75)
    add_text(slide, 0.42, bar_y, SLIDE_W_IN - 0.84, 0.46,
              "Core innovation: the agent performs Dynamic Tool Selection instead of running a fixed pipeline.",
              size=9.5, bold=True, italic=True, color=BLUE_DK, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)

    footer(slide, 2, question="Why is this more than a normal fraud classifier?")
    return slide


# ============================================================== SLIDE 3 ====
# HOW DOES IT WORK?  -- agentic system map with dynamic tool paths + gauge

def build_slide3(prs):
    slide = add_slide(prs)
    header_band(slide, 3, "Technical Approach", "Agentic System Map")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    cx0 = SLIDE_W_IN / 2

    # thin left rail with step numbers
    rail_x = 0.20
    add_line_shape(slide, rail_x, 1.02, rail_x, 6.95, color=BORDER_DK, weight=1.0)

    def rail_tick(y, n):
        add_oval(slide, rail_x - 0.05, y - 0.05, 0.10, fill=BLUE)

    # ---- 01 input chips ----
    in_y = 1.14
    rail_tick(in_y + 0.12, 1)
    inputs = [("✉", "SMS"), ("@", "EMAIL"), ("\U0001F517", "URL"), ("☎", "PHONE"),
              ("\U0001F4B3", "TRANSACTION")]
    n_in = len(inputs)
    chip_w = 1.62
    chip_gap = (content_w - chip_w * n_in) / (n_in - 1)
    ix = left_edge
    in_centers = []
    for icon, label in inputs:
        add_rect(slide, ix, in_y, chip_w, 0.36, fill=NAVY, radius=0.5)
        add_oval(slide, ix + 0.05, in_y + 0.05, 0.26, fill=NAVY_SOFT)
        add_text(slide, ix + 0.05, in_y + 0.05, 0.26, 0.26, icon, size=10, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, ix + 0.36, in_y, chip_w - 0.4, 0.36, label, size=8.6, bold=True,
                  color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        in_centers.append(ix + chip_w / 2)
        ix += chip_w + chip_gap
    add_text(slide, left_edge, in_y - 0.20, 3.0, 0.18, "01  INPUT SIGNALS", size=7.6, bold=True,
              color=TEXT_SOFT, spacing_pt=0.8)

    # ---- 02 orchestrator ----
    orch_r = 0.72
    orch_cy = in_y + 0.36 + 0.45 + orch_r
    for icx in in_centers:
        add_line_shape(slide, icx, in_y + 0.36, cx0, orch_cy - orch_r, color=BORDER_DK, weight=0.85)
    add_oval(slide, cx0 - orch_r - 0.10, orch_cy - orch_r - 0.10, (orch_r + 0.10) * 2, fill=NAVY_MED)
    add_oval(slide, cx0 - orch_r, orch_cy - orch_r, orch_r * 2, fill=TEAL_DK)
    add_text(slide, cx0 - orch_r, orch_cy - orch_r + 0.11, orch_r * 2, 0.4,
              "FRAUD INVESTIGATION\nAGENT", size=9.5, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, line_spacing=0.95)
    add_text(slide, cx0 - orch_r + 0.08, orch_cy + 0.0, orch_r * 2 - 0.16, 0.6,
              "Understand · Plan · Select Tools ·\nInvestigate · Fuse Evidence", size=6.8,
              color=TEAL_LIGHT, align=PP_ALIGN.CENTER, line_spacing=1.05)
    add_text(slide, cx0 - orch_r + 0.08, orch_cy + orch_r - 0.20, orch_r * 2 - 0.16, 0.18,
              "LangGraph + Gemini", size=6.6, italic=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, left_edge, orch_cy - orch_r - 0.30, 4.0, 0.18, "02  ORCHESTRATOR", size=7.6,
              bold=True, color=TEXT_SOFT, spacing_pt=0.8)

    # ---- 03 tool ecosystem with two highlighted dynamic paths ----
    tool_y = orch_cy + orch_r + 0.55
    tools = [
        ("Message\nAnalyzer", BLUE),
        ("URL\nAnalyzer", BLUE),
        ("Transaction\nAnalyzer", TEAL),
        ("Scam Pattern\nSearch", NAVY_SOFT),
        ("Threat\nIntelligence", NAVY_SOFT),
    ]
    n_t = len(tools)
    t_r = 0.40
    t_gap = (content_w - t_r * 2 * n_t) / (n_t - 1)
    tx = left_edge + t_r
    tool_centers = []
    for label, col in tools:
        highlighted = col in (BLUE, TEAL)
        line_col = col if highlighted else BORDER_DK
        line_w = 1.5 if highlighted else 0.85
        connect_circle_to_circle(slide, cx0, orch_cy, orch_r, tx, tool_y, t_r, color=line_col,
                                   weight=line_w)
        add_oval(slide, tx - t_r, tool_y - t_r, t_r * 2, fill=col if highlighted else CARD,
                 line=None if highlighted else BORDER_DK, line_w=1.0)
        tcolor = WHITE if highlighted else NAVY_DARK
        tf_shp = add_text(slide, tx - t_r + 0.04, tool_y - t_r, t_r * 2 - 0.08, t_r * 2, label,
                           size=7.3, bold=True, color=tcolor, align=PP_ALIGN.CENTER,
                           anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        tool_centers.append((tx, tool_y, highlighted, col))
        tx += t_r * 2 + t_gap
    lbl3_y = tool_y - t_r - 0.28
    add_rect(slide, left_edge - 0.04, lbl3_y - 0.02, 5.7, 0.22, fill=BG, radius=0,
             shape_type=MSO_SHAPE.RECTANGLE)
    add_text(slide, left_edge, lbl3_y, 5.7, 0.18,
              "03  TOOL ECOSYSTEM — agent decides which tools to invoke",
              size=7.6, bold=True, color=TEXT_SOFT, spacing_pt=0.6)
    leg_y = tool_y - t_r - 0.26
    add_oval(slide, SLIDE_W_IN - 0.42 - 3.05, leg_y + 0.02, 0.11, fill=BLUE)
    add_text(slide, SLIDE_W_IN - 0.42 - 2.88, leg_y - 0.02, 1.55, 0.18, "SMS + URL case",
              size=7.4, bold=True, color=BLUE_DK)
    add_oval(slide, SLIDE_W_IN - 0.42 - 1.35, leg_y + 0.02, 0.11, fill=TEAL)
    add_text(slide, SLIDE_W_IN - 0.42 - 1.18, leg_y - 0.02, 1.55, 0.18, "Transaction case",
              size=7.4, bold=True, color=TEAL_DK)

    # ---- 04 evidence fusion ----
    fus_cy = tool_y + t_r + 0.46
    fus_r = 0.34
    for (tcx, tcy, hl, col) in tool_centers:
        connect_circle_to_circle(slide, tcx, tcy, t_r, cx0, fus_cy, fus_r,
                                   color=(col if hl else BORDER_DK), weight=(1.3 if hl else 0.85))
    add_oval(slide, cx0 - fus_r, fus_cy - fus_r, fus_r * 2, fill=NAVY_DARK)
    add_text(slide, cx0 - fus_r, fus_cy - fus_r, fus_r * 2, fus_r * 2, "\U0001F517", size=13,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, cx0 - 1.3, fus_cy + fus_r + 0.04, 2.6, 0.18, "EVIDENCE FUSION", size=8.4,
              bold=True, color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.4)
    add_text(slide, left_edge, fus_cy - 0.09, 3.0, 0.18, "04  FUSE", size=7.6, bold=True,
              color=TEXT_SOFT, spacing_pt=0.8)

    # ---- 05 risk engine + gauge ----
    risk_y = fus_cy + fus_r + 0.26
    add_line_shape(slide, cx0, fus_cy + fus_r + 0.24, cx0, risk_y - 0.02, color=BORDER_DK, weight=1.0)
    gauge_w, gauge_h = 5.6, 0.30
    gx = cx0 - gauge_w / 2
    seg_w = gauge_w / 3
    add_rect(slide, gx, risk_y, seg_w, gauge_h, fill=GREEN, radius=0.5)
    add_rect(slide, gx + seg_w, risk_y, seg_w, gauge_h, fill=AMBER,
             shape_type=MSO_SHAPE.RECTANGLE)
    add_rect(slide, gx + 2 * seg_w, risk_y, seg_w, gauge_h, fill=RED, radius=0.5)
    for i, lbl in enumerate(["LOW", "MEDIUM", "HIGH"]):
        add_text(slide, gx + i * seg_w, risk_y + gauge_h + 0.03, seg_w, 0.18, lbl, size=7.6,
                  bold=True, color=TEXT_MUTED, align=PP_ALIGN.CENTER, spacing_pt=0.5)
    add_text(slide, gx - 2.1, risk_y - 0.02, 1.9, gauge_h, "ML + RULES\n+ EVIDENCE  →", size=8,
              bold=True, color=NAVY_DARK, align=PP_ALIGN.RIGHT, line_spacing=0.95,
              anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, gx + gauge_w + 0.15, risk_y - 0.02, 1.7, gauge_h, "EXPLAINABLE\nRISK SCORE",
              size=8, bold=True, color=NAVY_DARK, line_spacing=0.95, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, left_edge, risk_y - 0.20, 3.0, 0.18, "05  SCORE", size=7.6, bold=True,
              color=TEXT_SOFT, spacing_pt=0.8)

    # ---- 06 decision + 07 platform ----
    dec_y = risk_y + gauge_h + 0.32
    decisions = [("\U0001F4AC", "EXPLAIN"), ("✅", "RECOMMEND"), ("\U0001F441", "REVIEW")]
    n_d = len(decisions)
    d_w = 2.4
    d_gap = (content_w - d_w * n_d) / (n_d - 1)
    dxp = left_edge + (content_w - (d_w * n_d + d_gap * (n_d - 1))) / 2
    for icon, label in decisions:
        add_rect(slide, dxp, dec_y, d_w, 0.38, fill=BLUE_PALE, radius=0.5, line=BLUE_LIGHT,
                 line_w=0.75)
        add_text(slide, dxp, dec_y, 0.5, 0.38, icon, size=12, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, dxp + 0.42, dec_y, d_w - 0.5, 0.38, label, size=9, bold=True,
                  color=BLUE_DK, anchor=MSO_ANCHOR.MIDDLE)
        dxp += d_w + d_gap
    add_text(slide, left_edge, dec_y - 0.20, 3.0, 0.18, "06  DECIDE", size=7.6, bold=True,
              color=TEXT_SOFT, spacing_pt=0.8)

    plat_y = dec_y + 0.38 + 0.20
    platform = ["Supabase", "Cases", "Alerts", "Analytics", "Feedback"]
    n_p = len(platform)
    p_w = 1.9
    p_gap = (content_w - p_w * n_p) / (n_p - 1)
    pxp = left_edge
    for label in platform:
        add_rect(slide, pxp, plat_y, p_w, 0.32, fill=NAVY_DARK, radius=0.5)
        add_text(slide, pxp, plat_y, p_w, 0.32, label, size=8.2, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        pxp += p_w + p_gap
    add_text(slide, left_edge, plat_y - 0.20, 3.0, 0.18, "07  STORE & TRACK", size=7.6, bold=True,
              color=TEXT_SOFT, spacing_pt=0.8)

    footer(slide, 3, question="This is what makes it an agentic system.")
    return slide


# ============================================================== SLIDE 4 ====
# CAN WE BUILD IT?  -- build->validate->scale spine + risk-control map

def build_slide4(prs):
    slide = add_slide(prs)
    header_band(slide, 4, "Feasibility and Viability", "Build → Validate → Scale")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.00

    # ---- stage spine (pentagon pointer blocks) ----
    stages = [
        ("BUILD", NAVY_SOFT, "React · Supabase · LangGraph ·\nGemini · Scikit-learn"),
        ("VALIDATE", TEAL, "Datasets · Model Evaluation ·\nHuman Review · Feedback Loop"),
        ("SCALE", BLUE, "Threat Intel · More Fraud Types ·\nMore Tools · Cloud Deployment"),
    ]
    n_s = len(stages)
    arrow_gap = 0.34
    block_w = (content_w - arrow_gap * (n_s - 1)) / n_s
    block_h = 1.9
    bx = left_edge
    for i, (title, col, chips) in enumerate(stages):
        shp = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, Inches(bx), Inches(top), Inches(block_w),
                                       Inches(block_h))
        set_fill(shp, col)
        no_line(shp)
        no_shadow(shp)
        add_soft_shadow(shp, blur_pt=10, alpha_pct=18)
        safe_w = block_w * 0.66
        add_text(slide, bx + 0.22, top + 0.20, safe_w, 0.32, f"0{i+1}", size=11, bold=True,
                  color=WHITE, spacing_pt=0.5)
        add_text(slide, bx + 0.22, top + 0.52, safe_w, 0.4, title, size=17, bold=True, color=WHITE,
                  spacing_pt=0.6)
        add_text(slide, bx + 0.22, top + 0.98, safe_w, 0.8, chips, size=8.6, color=WHITE,
                  line_spacing=1.25)
        if i < n_s - 1:
            right_arrow(slide, bx + block_w + 0.03, top + block_h / 2, size=0.24, color=BORDER_DK)
        bx += block_w + arrow_gap

    # ---- risk-control map ----
    map_y = top + block_h + 0.24
    section_label(slide, left_edge, map_y, "Challenge → Mitigation Control Map", accent=RED)
    rows_y = map_y + 0.34
    pairs = [
        ("False positives", "Human review + feedback loop + threshold tuning"),
        ("AI hallucination", "LLM explains structured evidence — never invents it"),
        ("Threat-intel limitations", "Fallback local analysis + configurable adapters"),
        ("API / rate limits", "Caching + graceful fallback"),
        ("Sensitive user data", "Auth + access control + minimal collection"),
    ]
    row_h = 0.44
    mid_x = SLIDE_W_IN / 2
    for i, (chal, mit) in enumerate(pairs):
        ry = rows_y + i * row_h
        if i > 0:
            add_line_shape(slide, left_edge, ry, SLIDE_W_IN - left_edge, ry, color=BORDER,
                            weight=0.6)
        cy = ry + row_h / 2 - 0.14
        add_oval(slide, left_edge, cy, 0.28, fill=RED_LIGHT, line=RED, line_w=1.0)
        add_text(slide, left_edge, cy, 0.28, 0.28, "!", size=12, bold=True, color=RED,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, left_edge + 0.40, cy - 0.03, mid_x - left_edge - 0.9, 0.34, chal,
                  size=9.6, bold=True, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)
        right_arrow(slide, mid_x - 0.42, ry + row_h / 2, size=0.22, color=BORDER_DK)
        add_oval(slide, mid_x + 0.05, cy, 0.28, fill=GREEN_LIGHT, line=GREEN, line_w=1.0)
        add_text(slide, mid_x + 0.05, cy, 0.28, 0.28, "✓", size=11, bold=True, color=GREEN,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, mid_x + 0.45, cy - 0.03, SLIDE_W_IN - left_edge - mid_x - 0.5, 0.34, mit,
                  size=9.2, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)

    # ---- why buildable checklist ----
    chk_y = rows_y + len(pairs) * row_h + 0.20
    section_label(slide, left_edge, chk_y, "Why It Is Buildable", accent=GREEN)
    strip_y = chk_y + 0.32
    checks = ["Open-source ML", "Serverless backend", "Modular agent tools",
              "Free-tier prototype stack", "Incremental deployment"]
    n_c = len(checks)
    c_gap = 0.14
    c_w = (content_w - c_gap * (n_c - 1)) / n_c
    cx = left_edge
    for chk in checks:
        add_rect(slide, cx, strip_y, c_w, 0.42, fill=GREEN_LIGHT, radius=0.5)
        add_oval(slide, cx + 0.08, strip_y + 0.08, 0.26, fill=GREEN)
        add_text(slide, cx + 0.08, strip_y + 0.08, 0.26, 0.26, "✓", size=10, bold=True,
                  color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx + 0.40, strip_y, c_w - 0.48, 0.42, chk, size=8.0, bold=True,
                  color="0F5A28", anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        cx += c_w + c_gap

    footer(slide, 4, question="Yes — with a clear build path and risk controls.")
    return slide


# ============================================================== SLIDE 5 ====
# WHY DOES IT MATTER?  -- ecosystem cross + before/after with mock UI

def build_slide5(prs):
    slide = add_slide(prs)
    header_band(slide, 5, "Impact and Benefits", "Who It Serves & What Changes")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.00

    # ---- ecosystem cross ----
    cx0 = SLIDE_W_IN / 2
    hub_d = 0.72
    node_w, node_h = 3.05, 0.88
    v_gap = 0.16
    cy0 = top + node_h + v_gap + hub_d / 2

    add_oval(slide, cx0 - hub_d / 2 - 0.09, cy0 - hub_d / 2 - 0.09, hub_d + 0.18, fill=NAVY_MED)
    add_shield(slide, cx0 - hub_d * 0.36, cy0 - hub_d * 0.40, hub_d * 0.72, hub_d * 0.72,
               fill=BLUE, check=True, check_color=WHITE, check_size=hub_d * 0.72)
    add_text(slide, cx0 - 1.0, cy0 + hub_d / 2 + 0.05, 2.0, 0.18, "NETRAX", size=8.5, bold=True,
              color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.6)

    def cross_node(cxp, cyp, icon, title, bullets, col, from_dir):
        rx, ry = cxp - node_w / 2, cyp - node_h / 2
        add_rect(slide, rx, ry, node_w, node_h, fill=CARD, radius=0.10, line=BORDER, line_w=0.75,
                 shadow=True)
        add_rect(slide, rx, ry, 0.055, node_h, fill=col, shape_type=MSO_SHAPE.RECTANGLE)
        add_oval(slide, rx + 0.14, ry + 0.14, 0.34, fill=col)
        add_text(slide, rx + 0.14, ry + 0.14, 0.34, 0.34, icon, size=13, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, rx + 0.56, ry + 0.10, node_w - 0.7, 0.30, title, size=9.6, bold=True,
                  color=NAVY_DARK, line_spacing=0.95)
        add_bullets(slide, rx + 0.18, ry + 0.54, node_w - 0.34, node_h - 0.58, bullets, size=7.4,
                    bullet_color=col, space_after=1, line_spacing=1.0)
        connect_circle_to_rect(slide, cx0, cy0, hub_d / 2, cxp, cyp, node_w / 2, node_h / 2,
                                color=BORDER_DK, weight=1.1)

    cross_node(cx0, top + node_h / 2, "\U0001F464", "INDIVIDUAL USERS",
               ["Faster clarity on suspicious content", "Clear protective guidance"], NAVY, "N")
    bottom_cy = cy0 + hub_d / 2 + v_gap + node_h / 2
    cross_node(cx0, bottom_cy, "\U0001F3E2", "ENTERPRISES",
               ["Reduced manual screening effort", "Centralized case history"], NAVY_SOFT, "S")
    side_w = 3.35
    left_cx = left_edge + side_w / 2
    cross_node(left_cx, cy0, "\U0001F6E1", "CYBERSECURITY TEAMS",
               ["Evidence-based investigation", "Tool execution trace"], TEAL, "W")
    right_cx = SLIDE_W_IN - left_edge - side_w / 2
    cross_node(right_cx, cy0, "\U0001F3E6", "FINANCIAL / SERVICE\nORGANIZATIONS",
               ["Structured fraud investigation", "Explainable risk signals"], BLUE, "E")

    zone1_bottom = bottom_cy + node_h / 2
    add_text(slide, left_edge, zone1_bottom + 0.02, content_w, 0.2,
              "A decision-support & early-warning platform — not a replacement for banks, police or official cybercrime systems.",
              size=8, italic=True, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    # ---- BEFORE / mock UI AFTER ----
    z2_y = zone1_bottom + 0.22
    z2_h = 2.30
    bw = 3.15
    add_rect(slide, left_edge, z2_y, bw, z2_h, fill=RED_LIGHT, radius=0.07, line="F3B9B9",
             line_w=0.75)
    add_text(slide, left_edge + 0.18, z2_y + 0.12, bw - 0.36, 0.20, "BEFORE NETRAX", size=9.2,
              bold=True, color=RED_DK, spacing_pt=0.3)
    steps_b = ["Suspicious message", "Confusion", "Delayed response", "Potential exposure"]
    sb_h = 0.28
    sb_y = z2_y + 0.40
    for i, s in enumerate(steps_b):
        add_rect(slide, left_edge + 0.22, sb_y, bw - 0.44, sb_h, fill=WHITE, radius=0.5, line=RED,
                 line_w=0.9)
        add_text(slide, left_edge + 0.22, sb_y, bw - 0.44, sb_h, s, size=8.2, bold=True,
                  color=RED_DK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps_b) - 1:
            down_arrow(slide, left_edge + bw / 2, sb_y + sb_h + 0.015, size=0.10, color=RED)
        sb_y += sb_h + 0.115

    # ---- mock UI card (illustrative) ----
    mx = left_edge + bw + 0.26
    mw = SLIDE_W_IN - left_edge - mx
    add_rect(slide, mx, z2_y, mw, z2_h, fill=CARD, radius=0.07, line=BORDER_DK, line_w=1.0,
             shadow=True)
    add_rect(slide, mx, z2_y, mw, 0.40, fill=NAVY_DARK, radius=0.07, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(slide, mx, z2_y + 0.20, mw, 0.20, fill=NAVY_DARK, shape_type=MSO_SHAPE.RECTANGLE)
    add_text(slide, mx + 0.18, z2_y, 2.6, 0.40, "CASE #FS-1024", size=10.5, bold=True, color=WHITE,
              anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, mx + 2.0, z2_y, 2.6, 0.40, "Input: Suspicious SMS", size=8.6, italic=True,
              color=BLUE_LIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, mx + mw - 1.9, z2_y + 0.09, 1.72, 0.22, fill=AMBER, radius=0.5)
    add_text(slide, mx + mw - 1.9, z2_y + 0.09, 1.72, 0.22, "SAMPLE · ILLUSTRATIVE UI", size=6.8,
              bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    body_y = z2_y + 0.50
    col1_w = mw * 0.40
    add_text(slide, mx + 0.18, body_y, col1_w - 0.2, 0.2, "AI INVESTIGATION", size=8.2, bold=True,
              color=NAVY_DARK, spacing_pt=0.4)
    checks = ["Input classified", "URL extracted", "Message analyzed", "URL analyzed",
              "Scam pattern matched", "Evidence fused"]
    for i, ch in enumerate(checks):
        cy = body_y + 0.26 + i * 0.235
        add_oval(slide, mx + 0.18, cy, 0.15, fill=GREEN)
        add_text(slide, mx + 0.18, cy - 0.015, 0.15, 0.15, "✓", size=7, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, mx + 0.40, cy - 0.05, col1_w - 0.5, 0.22, ch, size=8.0, color=TEXT_DARK)

    col2_x = mx + col1_w + 0.10
    col2_w = mw * 0.24
    badge_d = 1.05
    bxp = col2_x + (col2_w - badge_d) / 2
    add_oval(slide, bxp, body_y + 0.10, badge_d, fill=RED)
    add_text(slide, bxp, body_y + 0.24, badge_d, 0.4, "94%", size=17, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER)
    add_text(slide, bxp, body_y + 0.62, badge_d, 0.2, "HIGH RISK", size=7.6, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, spacing_pt=0.4)
    add_text(slide, col2_x, body_y + badge_d + 0.16, col2_w, 0.2, "RISK SCORE", size=7.2, bold=True,
              color=TEXT_SOFT, align=PP_ALIGN.CENTER, spacing_pt=0.5)

    col3_x = col2_x + col2_w + 0.14
    col3_w = mx + mw - col3_x - 0.16
    add_text(slide, col3_x, body_y, col3_w, 0.2, "WHY", size=8.2, bold=True, color=NAVY_DARK,
              spacing_pt=0.4)
    whys = ["Suspicious URL", "Urgency language", "Financial request"]
    for i, w in enumerate(whys):
        wy = body_y + 0.23 + i * 0.225
        add_rect(slide, col3_x, wy, col3_w, 0.19, fill=RED_LIGHT, radius=0.5)
        add_text(slide, col3_x, wy, col3_w, 0.19, w, size=7.2, bold=True, color=RED_DK,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, col3_x, body_y + 0.92, col3_w, 0.2, "ACTION", size=8.2, bold=True,
              color=NAVY_DARK, spacing_pt=0.4)
    acts = ["Do not click", "Do not share OTP", "Verify officially"]
    for i, a in enumerate(acts):
        ay = body_y + 1.15 + i * 0.205
        add_oval(slide, col3_x, ay, 0.14, fill=GREEN)
        add_text(slide, col3_x, ay - 0.02, 0.14, 0.14, "✓", size=6.6, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, col3_x + 0.20, ay - 0.05, col3_w - 0.2, 0.2, a, size=7.6, color=TEXT_DARK)

    # ---- impact pillars ----
    p_y = z2_y + z2_h + 0.18
    pillars = [
        ("SOCIAL", "Digital safety awareness", GREEN),
        ("SECURITY", "Evidence-based investigation", NAVY_SOFT),
        ("ECONOMIC", "Reduced avoidable exposure; streamlined initial screening", BLUE),
    ]
    gap = 0.22
    pw = (content_w - gap * 2) / 3
    ph = 0.58
    cx = left_edge
    for title, desc, col in pillars:
        add_rect(slide, cx, p_y, pw, ph, fill=CARD, radius=0.5, line=BORDER, line_w=0.75)
        add_oval(slide, cx + 0.12, p_y + ph / 2 - 0.12, 0.24, fill=col)
        add_text(slide, cx + 0.46, p_y + 0.06, pw - 0.6, 0.2, title, size=8.6, bold=True, color=col,
                  spacing_pt=0.4)
        add_text(slide, cx + 0.46, p_y + 0.26, pw - 0.6, 0.28, desc, size=7.6, color=TEXT_MUTED,
                  line_spacing=1.0)
        cx += pw + gap

    footer(slide, 5, question="This solves a real user problem.")
    return slide


# ============================================================== SLIDE 6 ====
# WHAT SUPPORTS IT?  -- evidence board

def build_slide6(prs):
    slide = add_slide(prs)
    header_band(slide, 6, "Research and References", "Evidence Board")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.02

    clusters = [
        ("\U0001F4CA", "DATA", BLUE, [
            ("UCI SMS Spam Collection", "archive.ics.uci.edu/dataset/228"),
            ("UCI Phishing Websites", "archive.ics.uci.edu/dataset/327"),
            ("Credit Card Fraud Detection", "Kaggle / ULB dataset"),
        ]),
        ("\U0001F6E1", "INTELLIGENCE", TEAL, [
            ("PhishTank", "phishtank.org"),
            ("Threat Intelligence Feeds", "configurable adapters (planned)"),
        ]),
        ("⚙", "TECHNOLOGY", AMBER, [
            ("LangGraph", "langchain-ai.github.io/langgraph"),
            ("Supabase", "supabase.com/docs"),
            ("Gemini API", "ai.google.dev/gemini-api/docs"),
            ("Scikit-learn", "scikit-learn.org"),
        ]),
    ]
    n_c = len(clusters)
    gap = 0.24
    col_w = (content_w - gap * (n_c - 1)) / n_c
    col_h = 3.0
    cx = left_edge
    for icon, title, col, items in clusters:
        add_rect(slide, cx, top, col_w, col_h, fill=CARD, radius=0.08, line=BORDER, line_w=0.75,
                 shadow=True)
        add_rect(slide, cx, top, col_w, 0.06, fill=col, shape_type=MSO_SHAPE.RECTANGLE)
        add_oval(slide, cx + 0.20, top + 0.24, 0.48, fill=col)
        add_text(slide, cx + 0.20, top + 0.24, 0.48, 0.48, icon, size=18, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx + 0.80, top + 0.32, col_w - 0.96, 0.3, title, size=13, bold=True,
                  color=NAVY_DARK, spacing_pt=0.6)
        iy = top + 0.98
        for name, sub in items:
            risk_dot(slide, cx + 0.22, iy + 0.06, 0.10, col)
            add_text(slide, cx + 0.42, iy, col_w - 0.60, 0.22, name, size=9.4, bold=True,
                      color=TEXT_DARK)
            add_text(slide, cx + 0.42, iy + 0.225, col_w - 0.60, 0.20, sub, size=7.8,
                      color=TEXT_MUTED, italic=True)
            iy += 0.50
        cx += col_w + gap

    # ---- security references strip ----
    sec_y = top + col_h + 0.22
    section_label(slide, left_edge, sec_y, "Security References", accent=NAVY_SOFT)
    sec_row_y = sec_y + 0.32
    sec_refs = ["NIST Cybersecurity Guidance", "OWASP", "CERT-In",
                "Govt. of India Cybercrime Awareness"]
    n_sr = len(sec_refs)
    sr_gap = 0.18
    sr_w = (content_w - sr_gap * (n_sr - 1)) / n_sr
    sx = left_edge
    for r in sec_refs:
        add_rect(slide, sx, sec_row_y, sr_w, 0.40, fill=BLUE_PALE, radius=0.08, line=BLUE_LIGHT,
                 line_w=0.75)
        add_text(slide, sx + 0.12, sec_row_y, sr_w - 0.24, 0.40, r, size=8.6, bold=True,
                  color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        sx += sr_w + sr_gap

    # ---- research basis tag row ----
    rb_y = sec_row_y + 0.40 + 0.22
    section_label(slide, left_edge, rb_y, "Research Basis", accent=TEAL)
    tag_y = rb_y + 0.32
    tags = ["Phishing Detection", "SMS Spam Detection", "Fraud Detection", "Explainable AI",
            "Agentic AI / Tool-Use"]
    tx = left_edge
    for t in tags:
        tw = 0.18 + len(t) * 0.082
        add_rect(slide, tx, tag_y, tw, 0.32, fill=TEAL_LIGHT, radius=0.5)
        add_text(slide, tx, tag_y, tw, 0.32, t, size=8.2, bold=True, color=TEAL_DK,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tx += tw + 0.14
    add_text(slide, tx + 0.06, tag_y + 0.02, 3.0, 0.28, "— reference to be finalized", size=7.6,
              italic=True, color=TEXT_SOFT, anchor=MSO_ANCHOR.MIDDLE)

    # ---- project links ----
    lk_y = tag_y + 0.32 + 0.22
    links = [("PROTOTYPE / DEMO", "[ADD WORKING PROTOTYPE LINK]"),
             ("GITHUB", "[ADD GITHUB LINK]"),
             ("DEMO VIDEO", "[ADD VIDEO LINK]")]
    n_l = len(links)
    l_gap = 0.22
    l_w = (content_w - l_gap * (n_l - 1)) / n_l
    lx = left_edge
    for lbl, val in links:
        add_rect(slide, lx, lk_y, l_w, 0.50, fill=NAVY_DARK, radius=0.10)
        add_text(slide, lx + 0.16, lk_y + 0.07, l_w - 0.32, 0.16, lbl, size=7.8, bold=True,
                  color=TEAL_LIGHT, spacing_pt=0.4)
        add_text(slide, lx + 0.16, lk_y + 0.25, l_w - 0.32, 0.20, val, size=9.2, bold=True,
                  color=WHITE)
        lx += l_w + l_gap

    footer(slide, 6, question="Grounded in real datasets, tools and standards.")
    return slide


def save_and_report(prs, path):
    prs.save(path)
    print("Saved:", path)


def main():
    prs = new_presentation()
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    build_slide4(prs)
    build_slide5(prs)
    build_slide6(prs)
    save_and_report(prs, r"c:\Surakshit Bharat\ppt\NetraX_SIH_2026_FINAL.pptx")


if __name__ == "__main__":
    main()
