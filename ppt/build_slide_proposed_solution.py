# -*- coding: utf-8 -*-
"""Single slide matching the official SIH template screenshot exactly:
oval team-name badge, idea-title header, SIH badge, blue underlined
subheading, bullet body, blue footer banner with page number."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = "1F3864"
BLUE = "1F4E9C"
BLUE_BANNER = "1544A0"
PURPLE = "6B4C9A"
BLACK = "0D0D0D"
TEXT = "1A1A1A"
WHITE = "FFFFFF"
SERIF = "Times New Roman"
SANS = "Calibri"
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5

TEAM_NAME = "Nexora"


def C(h):
    return RGBColor.from_string(h)


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None, line_w=1.25):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=12, color=TEXT, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             underline=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font; r.font.color.rgb = C(color); r.font.underline = underline
    return tb


def add_bullets(slide, x, y, w, h, items, size=15, color=TEXT, font=SANS, line_spacing=1.2, space_after=12):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r0 = p.add_run(); r0.text = "•  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = font; r0.font.color.rgb = C(color)
        r1 = p.add_run(); r1.text = item
        r1.font.size = Pt(size); r1.font.name = font; r1.font.color.rgb = C(color)
    return tb


prs = Presentation()
prs.slide_width = Inches(SLIDE_W_IN)
prs.slide_height = Inches(SLIDE_H_IN)
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W_IN, SLIDE_H_IN, fill=WHITE)

# team name oval, top-left
add_shape(slide, MSO_SHAPE.OVAL, 0.35, 0.22, 1.15, 0.62, fill=WHITE, line=PURPLE, line_w=1.5)
add_text(slide, 0.35, 0.22, 1.15, 0.62, TEAM_NAME, size=12, color=TEXT, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)

# idea title, top-center
add_text(slide, 1.7, 0.22, 9.0, 0.55, "NETRAX", size=27, bold=True, color=BLACK,
          align=PP_ALIGN.CENTER, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)

# SIH badge, top-right
hexs = add_shape(slide, MSO_SHAPE.HEXAGON, 10.9, 0.20, 0.62, 0.62, fill=NAVY)
add_text(slide, 10.9, 0.20, 0.62, 0.62, "SIH", size=13, bold=True, color=WHITE,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(slide, 11.62, 0.17, 2.1, 0.34, "SMART INDIA\nHACKATHON", size=10.5, bold=True, color=NAVY,
          line_spacing=0.95)
add_text(slide, 11.62, 0.52, 2.1, 0.22, "2026", size=10.5, bold=True, color=BLUE)

# subheading
add_text(slide, 0.55, 1.35, 0.30, 0.34, "❖", size=17, bold=True, color=BLUE)
add_text(slide, 0.87, 1.35, 11.9, 0.34, "Proposed Solution (Email Threat Detection, Investigation & Forensic Intelligence)",
          size=19, bold=True, color=BLUE, underline=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)

# bullets
add_bullets(slide, 0.55, 2.05, 12.2, 4.5, [
    "Detailed explanation: NetraX takes a suspicious email and runs it through header forensics, "
    "URL/domain analysis, threat-intelligence correlation (PhishTank, URLhaus), and IP/ASN "
    "geolocation — then fuses that evidence into one explainable, deterministic risk score with a "
    "recommended action.",
    "How it addresses the problem: traditional filters stop at “suspicious / not suspicious.” "
    "SOC teams and investigators are then left to manually piece together evidence across headers, "
    "links and infrastructure. NetraX automates that investigation and hands back cited evidence, "
    "not just a verdict.",
    "Innovation and uniqueness: an agentic orchestrator dynamically selects only the tools an "
    "email's own indicators call for — no URL means URL/threat-intel tools are skipped entirely, "
    "no public source IP means geolocation is skipped entirely — with every skip logged as an "
    "auditable event, never a silent no-op.",
], size=15, line_spacing=1.2, space_after=14)

# footer banner
y = SLIDE_H_IN - 0.42
add_shape(slide, MSO_SHAPE.RECTANGLE, 0, y, SLIDE_W_IN, 0.42, fill=BLUE_BANNER)
add_text(slide, 0, y, SLIDE_W_IN - 0.5, 0.42, "@SIH Idea submission- Template", size=11, color=WHITE,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(slide, SLIDE_W_IN - 0.9, y, 0.7, 0.42, "2", size=11, bold=True, color=WHITE,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save("NetraX_Proposed_Solution_Slide.pptx")
print("Saved NetraX_Proposed_Solution_Slide.pptx")
