# -*- coding: utf-8 -*-
"""NetraX - SIH26106 (Email Threat Detection, GeoLocation & Forensic
Intelligence Platform) deck. 6 slides, 16:9, python-pptx.

Content is grounded in the actual repo (README.md, docs/FINAL_STATUS.md,
docs/AGENT_ARCHITECTURE.md, docs/RISK_SCORING.md, docs/DATA_SOURCES.md,
docs/ml/TRAINING_REPORT.md) as of 2026-09-08. No invented metrics, no
invented team info beyond what the user supplied (Team Name: Nexora,
Team ID: not yet assigned).
"""
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- THEME ----
NAVY_DARK   = "0A1B33"
NAVY        = "10233E"
NAVY_MED    = "1C3A63"
NAVY_SOFT   = "27466F"
BLUE        = "2563EB"
BLUE_DK     = "1D4ED8"
BLUE_LIGHT  = "DBEAFE"
BLUE_PALE   = "EFF6FF"
TEAL        = "0D9488"
TEAL_DK     = "0F766E"
TEAL_LIGHT  = "CCFBF1"
INDIGO      = "4F46E5"
INDIGO_DK   = "4338CA"
INDIGO_LIGHT= "E0E7FF"
GREEN       = "16A34A"
GREEN_DK    = "15803D"
GREEN_LIGHT = "DCFCE7"
AMBER       = "D97706"
AMBER_LIGHT = "FEF3C7"
RED         = "DC2626"
RED_DK      = "991B1B"
RED_LIGHT   = "FEE2E2"
BG          = "F5F8FC"
CARD        = "FFFFFF"
BORDER      = "E2E8F0"
BORDER_DK   = "CBD5E1"
TEXT_DARK   = "0F172A"
TEXT_MUTED  = "55647E"
TEXT_SOFT   = "8592A8"
WHITE       = "FFFFFF"
FONT        = "Segoe UI"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

INFO = {
    "ps_id": "SIH26106",
    "ps_title": "AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform",
    "org": "All India Council for Technical Education (Cyber Security Cell)",
    "dept": "Cyber Security Cell",
    "theme": "Blockchain & Cybersecurity",
    "category": "Software",
    "team_id": "[TEAM ID]",
    "team_name": "Nexora",
}

# --------------------------------------------------------------- HELPERS ---

def C(hex_color):
    return RGBColor.from_string(hex_color)


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def add_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_fill(bg, BG)
    no_line(bg)
    bg.shadow.inherit = False
    return slide


def set_fill(shape, hex_color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(hex_color)


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, hex_color, width=0.75):
    shape.line.color.rgb = C(hex_color)
    shape.line.width = Pt(width)


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_soft_shadow(shape, blur_pt=9, dist_pt=2.2, dir_deg=90, alpha_pct=20, color="1E293B"):
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    effectLst = etree.SubElement(spPr, qn('a:effectLst'))
    outerShdw = etree.SubElement(effectLst, qn('a:outerShdw'))
    outerShdw.set('blurRad', str(int(blur_pt * 12700)))
    outerShdw.set('dist', str(int(dist_pt * 12700)))
    outerShdw.set('dir', str(int(dir_deg * 60000)))
    outerShdw.set('rotWithShape', '0')
    clr = etree.SubElement(outerShdw, qn('a:srgbClr'))
    clr.set('val', color)
    alpha = etree.SubElement(clr, qn('a:alpha'))
    alpha.set('val', str(int(alpha_pct * 1000)))


def add_rect(slide, x, y, w, h, fill=CARD, line=None, line_w=0.75, radius=0.07,
             shadow=False, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill)
    if line:
        set_line(shp, line, line_w)
    else:
        no_line(shp)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    no_shadow(shp)
    if shadow:
        add_soft_shadow(shp)
    shp.text_frame.margin_left = 0
    shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0
    shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, d, fill=BLUE, line=None, line_w=1.0, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    if fill is None:
        shp.fill.background()
    else:
        set_fill(shp, fill)
    if line:
        set_line(shp, line, line_w)
    else:
        no_line(shp)
    no_shadow(shp)
    if shadow:
        add_soft_shadow(shp)
    return shp


def add_line_shape(slide, x1, y1, x2, y2, color=BORDER, weight=1.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = C(color)
    conn.line.width = Pt(weight)
    if dash:
        d = conn.line._get_or_add_ln()
        pd = etree.SubElement(d, qn('a:prstDash'))
        pd.set('val', dash)
    no_shadow(conn)
    return conn


def add_text(slide, x, y, w, h, text, size=12, color=TEXT_DARK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             wrap=True, spacing_pt=0.0, caps=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text.upper() if caps else text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = C(color)
    if spacing_pt:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(int(spacing_pt * 100)))
    return tb


def add_multi(slide, x, y, w, h, runs, size=9, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              line_spacing=1.05, font=FONT):
    """runs: list of (text, color, bold, italic) tuples -> single paragraph, multiple runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for text, color, bold, italic in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = C(color)
    return tb


def add_bullets(slide, x, y, w, h, items, size=10.5, color=TEXT_DARK, bullet_color=BLUE,
                 line_spacing=1.1, space_after=5, font=FONT, anchor=MSO_ANCHOR.TOP,
                 bullet_char="•"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r0 = p.add_run()
        r0.text = bullet_char + "  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = font
        r0.font.color.rgb = C(bullet_color)
        r1 = p.add_run()
        r1.text = item
        r1.font.size = Pt(size); r1.font.name = font
        r1.font.color.rgb = C(color)
    return tb


def add_shield(slide, x, y, w, h, fill=BLUE, line=None, check=True, check_color=WHITE,
                check_size=None):
    pts = [(0.0, 0.10), (0.5, 0.0), (1.0, 0.10), (1.0, 0.50), (0.5, 1.0), (0.0, 0.50)]
    abs_pts = [(Inches(x + px * w), Inches(y + py * h)) for px, py in pts]
    fb = slide.shapes.build_freeform(abs_pts[0][0], abs_pts[0][1])
    fb.add_line_segments(abs_pts[1:], close=True)
    shp = fb.convert_to_shape()
    set_fill(shp, fill)
    if line:
        set_line(shp, line, 1.0)
    else:
        no_line(shp)
    no_shadow(shp)
    if check:
        cs = check_size or (h * 0.5)
        add_text(slide, x, y + h * 0.16, w, cs, "✓", size=int(cs * 62), color=check_color,
                  bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shp


def add_logo(slide, x, y, size=0.34, wordmark=True, on_dark=True, word_size=13):
    add_shield(slide, x, y, size, size, fill=BLUE if on_dark else NAVY, check=True,
               check_color=WHITE, check_size=size)
    if wordmark:
        tcol = WHITE if on_dark else NAVY_DARK
        add_text(slide, x + size + 0.11, y - 0.02, 3.2, size + 0.06, "NETRAX",
                  size=word_size, bold=True, color=tcol, anchor=MSO_ANCHOR.MIDDLE,
                  spacing_pt=0.3)


def header_band(slide, num, title, subtitle=None, height=0.86):
    add_rect(slide, 0, 0, SLIDE_W_IN, height, fill=NAVY_DARK, shape_type=MSO_SHAPE.RECTANGLE)
    add_rect(slide, 0, height, SLIDE_W_IN, 0.04, fill=BLUE, shape_type=MSO_SHAPE.RECTANGLE)
    add_logo(slide, 0.42, height / 2 - 0.15, size=0.30, wordmark=True, on_dark=True, word_size=13.5)
    add_oval(slide, 3.55, height / 2 - 0.17, 0.34, fill=BLUE)
    add_text(slide, 3.55, height / 2 - 0.17, 0.34, 0.34, str(num), size=13, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ty = 0.13 if subtitle else height / 2 - 0.20
    add_text(slide, 4.05, ty, 6.6, 0.44, title, size=21, bold=True, color=WHITE, caps=True,
              spacing_pt=0.4, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 4.05, ty + 0.40, 6.6, 0.28, subtitle, size=11, color=BLUE_LIGHT, italic=True)
    add_text(slide, 10.55, 0.14, 2.35, 0.22, "SMART INDIA HACKATHON", size=9, bold=True,
              color=TEAL_LIGHT, align=PP_ALIGN.RIGHT, spacing_pt=0.4)
    add_text(slide, 10.55, 0.37, 2.35, 0.28, "2026  ·  SOFTWARE", size=12, bold=True,
              color=WHITE, align=PP_ALIGN.RIGHT)


def footer(slide, page_num, total=6, question=""):
    y = 7.20
    add_line_shape(slide, 0.42, y, SLIDE_W_IN - 0.42, y, color=BORDER_DK, weight=0.75)
    add_shield(slide, 0.42, y + 0.075, 0.15, 0.15, fill=NAVY, check=True, check_color=WHITE,
               check_size=0.15)
    add_text(slide, 0.64, y + 0.05, 5.5, 0.22, "NETRAX  ·  Detect · Investigate · Explain · Act",
              size=8, color=TEXT_MUTED, anchor=MSO_ANCHOR.MIDDLE)
    if question:
        add_text(slide, 4.5, y + 0.05, 5.3, 0.22, question, size=8, italic=True, color=TEXT_SOFT,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 9.5, y + 0.05, 3.38, 0.22, f"Page {page_num} / {total}",
              size=8, color=TEXT_MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def right_arrow(slide, x, y_center, size=0.16, color=BORDER_DK):
    w = size
    h = size * 0.78
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y_center - h / 2),
                                  Inches(w), Inches(h))
    set_fill(shp, color); no_line(shp); no_shadow(shp)
    try:
        shp.adjustments[0] = 0.55; shp.adjustments[1] = 0.55
    except Exception:
        pass
    return shp


def down_arrow(slide, x_center, y, size=0.14, color=BORDER_DK):
    w = size * 1.3; h = size
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x_center - w / 2), Inches(y),
                                  Inches(w), Inches(h))
    set_fill(shp, color); no_line(shp); no_shadow(shp)
    try:
        shp.adjustments[0] = 0.55; shp.adjustments[1] = 0.55
    except Exception:
        pass
    return shp


def section_label(slide, x, y, text, color=NAVY_DARK, size=12.5, accent=BLUE):
    add_rect(slide, x, y + 0.03, 0.07, 0.20, fill=accent, radius=0.5)
    add_text(slide, x + 0.16, y, 8.5, 0.26, text, size=size, bold=True, color=color, caps=True,
              spacing_pt=0.3)


def circle_edge_point(cx, cy, r, dx, dy):
    n = math.hypot(dx, dy) or 1.0
    ux, uy = dx / n, dy / n
    return cx + ux * r, cy + uy * r


def rect_edge_point(cx, cy, hw, hh, dx, dy):
    if dx == 0 and dy == 0:
        return cx, cy
    tx = hw / abs(dx) if dx != 0 else float('inf')
    ty = hh / abs(dy) if dy != 0 else float('inf')
    t = min(tx, ty)
    return cx + dx * t, cy + dy * t


def connect_circle_to_circle(slide, c1x, c1y, r1, c2x, c2y, r2, color=BORDER_DK, weight=1.1,
                              dash=None):
    dx, dy = c2x - c1x, c2y - c1y
    x1, y1 = circle_edge_point(c1x, c1y, r1, dx, dy)
    x2, y2 = circle_edge_point(c2x, c2y, r2, -dx, -dy)
    add_line_shape(slide, x1, y1, x2, y2, color=color, weight=weight, dash=dash)


def add_circuit_pattern(slide, x0, y0, cols, rows, spacing, dot_d=0.035, color=NAVY_SOFT):
    centers = []
    for r in range(rows):
        for c in range(cols):
            cx = x0 + c * spacing
            cy = y0 + r * spacing
            centers.append((cx, cy))
            add_oval(slide, cx - dot_d / 2, cy - dot_d / 2, dot_d, fill=color)
    for r in range(rows):
        for c in range(cols - 1):
            i = r * cols + c
            x1, y1 = centers[i]
            x2, y2 = centers[i + 1]
            add_line_shape(slide, x1, y1, x2, y2, color=color, weight=0.75)


def node(slide, cx, cy, r, icon, label, fill=BLUE, icon_color=WHITE, label_color=NAVY_DARK,
         label_size=8, icon_size=None, dashed=False, label_below=True, label_w=1.7,
         label_bold=True):
    line = BORDER_DK if dashed else None
    add_oval(slide, cx - r, cy - r, r * 2, fill=fill, line=line, line_w=1.2 if dashed else 1.0)
    add_text(slide, cx - r, cy - r, r * 2, r * 2, icon, size=icon_size or int(r * 34),
              color=icon_color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if label_below:
        add_text(slide, cx - label_w / 2, cy + r + 0.04, label_w, 0.32, label, size=label_size,
                  bold=label_bold, color=label_color, align=PP_ALIGN.CENTER, line_spacing=0.95)
    return cx, cy, r


# ============================================================== SLIDE 1 ====

def build_slide1(prs):
    slide = add_slide(prs)
    add_rect(slide, 0, 0, SLIDE_W_IN, SLIDE_H_IN, fill=NAVY_DARK, shape_type=MSO_SHAPE.RECTANGLE)
    add_rect(slide, 9.9, -1.7, 6.0, 6.0, fill=NAVY, shape_type=MSO_SHAPE.OVAL)
    add_rect(slide, -2.2, 6.3, 5.0, 5.0, fill=NAVY, shape_type=MSO_SHAPE.OVAL)
    add_circuit_pattern(slide, 11.95, 0.30, 3, 3, 0.24, dot_d=0.032, color=NAVY_SOFT)

    add_text(slide, 0, 0.28, SLIDE_W_IN, 0.26, "SMART INDIA HACKATHON 2026", size=12.5, bold=True,
              color=TEAL_LIGHT, align=PP_ALIGN.CENTER, spacing_pt=2.6)

    # ---- hero investigation-story flow ----
    flow_y = 1.28
    stages = [
        ("✉", "SUSPICIOUS\nEMAIL", NAVY_SOFT),
        ("\U0001F50D", "AI\nINVESTIGATION", BLUE),
        ("\U0001F4CB", "FORENSIC\nEVIDENCE", TEAL),
        ("\U0001F310", "THREAT\nINFRASTRUCTURE", INDIGO),
        ("check", "RISK &\nACTION", GREEN),
    ]
    n = len(stages)
    r = 0.40
    margin = 2.05
    usable = SLIDE_W_IN - 2 * margin
    xs = [margin + usable * (i / (n - 1)) for i in range(n)]
    cy = flow_y + r
    for i in range(n - 1):
        add_line_shape(slide, xs[i] + r, cy, xs[i + 1] - r, cy, color="2E4E7E", weight=1.4)
        mx = (xs[i] + r + xs[i + 1] - r) / 2
        tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(mx - 0.05), Inches(cy - 0.05),
                                      Inches(0.10), Inches(0.10))
        tri.rotation = 90
        set_fill(tri, "3B5D8A"); no_line(tri); no_shadow(tri)
    for i, (icon, label, col) in enumerate(stages):
        cx = xs[i]
        add_oval(slide, cx - r - 0.09, cy - r - 0.09, (r + 0.09) * 2, fill=NAVY_MED)
        if icon == "check":
            add_shield(slide, cx - r * 0.75, cy - r * 0.82, r * 1.5, r * 1.5, fill=col,
                       check=True, check_color=WHITE, check_size=r * 1.5)
        else:
            add_oval(slide, cx - r, cy - r, r * 2, fill=col)
            add_text(slide, cx - r, cy - r, r * 2, r * 2, icon, size=17, color=WHITE,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx - 0.95, cy + r + 0.10, 1.9, 0.42, label, size=8.6, bold=True,
                  color=BLUE_LIGHT, align=PP_ALIGN.CENTER, line_spacing=0.98)

    add_text(slide, 0, 2.62, SLIDE_W_IN, 0.24,
              "“Others detect the threat. NetraX investigates the threat.”",
              size=12, italic=True, bold=True, color=TEAL_LIGHT, align=PP_ALIGN.CENTER)

    # ---- title block ----
    add_text(slide, 0, 3.10, SLIDE_W_IN, 0.62, "NETRAX", size=42, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, spacing_pt=1.6)
    add_text(slide, 1.4, 3.76, SLIDE_W_IN - 2.8, 0.55,
              "AI-Powered Email Threat Detection, GeoLocation & Forensic Intelligence Platform",
              size=14.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=1.08)
    add_text(slide, 1.7, 4.28, SLIDE_W_IN - 3.4, 0.28,
              "Detection tells you WHAT. NetraX investigates WHY.",
              size=11, italic=True, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)

    # ---- info band ----
    info_y, info_h = 4.72, 1.34
    add_rect(slide, 0.7, info_y, SLIDE_W_IN - 1.4, info_h, fill=NAVY_MED, radius=0.08,
             line="2C4B78", line_w=0.75)
    row1 = [
        ("PROBLEM STATEMENT ID", INFO["ps_id"]),
        ("THEME", INFO["theme"]),
        ("PS CATEGORY", INFO["category"]),
    ]
    row1b = [
        ("TEAM ID", INFO["team_id"]),
        ("TEAM NAME", INFO["team_name"]),
    ]
    col_w = (SLIDE_W_IN - 1.4 - 0.6) / 5
    col_x0 = 0.7 + 0.3
    ry = info_y + 0.16
    for ci, (lbl, val) in enumerate(row1):
        cxp = col_x0 + ci * col_w
        add_text(slide, cxp, ry, col_w - 0.1, 0.17, lbl, size=7.4, bold=True, color=TEAL_LIGHT, spacing_pt=0.7)
        add_text(slide, cxp, ry + 0.185, col_w - 0.1, 0.22, val, size=10.5, bold=True, color=WHITE)
    for ci, (lbl, val) in enumerate(row1b):
        cxp = col_x0 + (3 + ci) * col_w
        add_text(slide, cxp, ry, col_w - 0.1, 0.17, lbl, size=7.4, bold=True, color=TEAL_LIGHT, spacing_pt=0.7)
        add_text(slide, cxp, ry + 0.185, col_w - 0.1, 0.22, val, size=10.5, bold=True, color=WHITE)
    add_line_shape(slide, col_x0, ry + 0.52, SLIDE_W_IN - 0.7 - 0.3, ry + 0.52, color="2C4B78", weight=0.75)
    add_text(slide, col_x0, ry + 0.60, SLIDE_W_IN - 1.4 - 0.6, 0.2, "PROBLEM STATEMENT TITLE",
              size=7.4, bold=True, color=TEAL_LIGHT, spacing_pt=0.7)
    add_text(slide, col_x0, ry + 0.78, SLIDE_W_IN - 1.4 - 0.6, 0.24, INFO["ps_title"],
              size=10.5, bold=True, color=WHITE)

    add_text(slide, 0.7, 6.20, SLIDE_W_IN - 1.4, 0.22, INFO["org"], size=9.5, bold=True,
              color=BLUE_LIGHT, align=PP_ALIGN.CENTER)
    add_text(slide, 0.7, 6.44, SLIDE_W_IN - 1.4, 0.44,
              "Decision-support & forensic-intelligence prototype — not a substitute for banks, "
              "police or official cybercrime reporting (cybercrime.gov.in)  ·  SIH 2026",
              size=8.3, italic=True, color=TEXT_SOFT, align=PP_ALIGN.CENTER, line_spacing=1.15)
    return slide


# ============================================================== SLIDE 2 ====

def build_slide2(prs):
    slide = add_slide(prs)
    header_band(slide, 2, "Proposed Solution & Innovation", "From detecting a suspicious email to investigating it")
    left_edge, content_w = 0.42, SLIDE_W_IN - 0.84
    top = 1.06

    # ---- top: linear investigation chain ----
    chain = [
        ("✉", "SUSPICIOUS\nEMAIL", NAVY_SOFT),
        ("\U0001F4C4", "EMAIL\nFORENSICS", BLUE),
        ("\U0001F9E0", "AGENTIC\nINVESTIGATION", TEAL_DK),
        ("\U0001F6E1", "THREAT\nINTELLIGENCE", TEAL),
        ("\U0001F310", "INFRASTRUCTURE\nANALYSIS", INDIGO),
        ("\U0001F517", "EVIDENCE\nCORRELATION", NAVY_SOFT),
        ("⚖", "EXPLAINABLE\nRISK", AMBER),
    ]
    n = len(chain)
    r = 0.34
    margin = 0.85
    usable = content_w - 2 * (margin - left_edge)
    xs = [left_edge + margin - left_edge + usable * (i / (n - 1)) for i in range(n)]
    cy = top + r + 0.12
    for i in range(n - 1):
        add_line_shape(slide, xs[i] + r, cy, xs[i + 1] - r, cy, color=BORDER_DK, weight=1.3)
    centers = []
    for i, (icon, label, col) in enumerate(chain):
        cx = xs[i]
        hub = (i == 2)
        rr = r * 1.18 if hub else r
        add_oval(slide, cx - rr, cy - rr, rr * 2, fill=col)
        add_text(slide, cx - rr, cy - rr, rr * 2, rr * 2, icon, size=15 if hub else 13, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx - 0.75, cy + rr + 0.06, 1.5, 0.34, label, size=7.4, bold=True,
                  color=NAVY_DARK, align=PP_ALIGN.CENTER, line_spacing=0.92)
        centers.append((cx, cy, rr, col))

    # ---- innovation evidence-board: 6 leader-lined annotations ----
    label_bottom = cy + r + 0.10 + 0.34  # below each chain node's 2-line label
    board_top = label_bottom + 0.34
    innov = [
        (2, "\U0001F9E0", "AGENTIC INVESTIGATION", "AI dynamically selects only the tools an email's own indicators call for — no fixed pipeline.", TEAL_DK),
        (1, "\U0001F4C4", "EMAIL FORENSICS", "Sender identity, SPF/DKIM/DMARC, Received-chain and header anomalies, fully cited.", BLUE),
        (3, "\U0001F6E1", "THREAT INTELLIGENCE", "Extracted URLs/domains correlated against PhishTank and URLhaus where applicable.", TEAL),
        (4, "\U0001F310", "INFRASTRUCTURE INTELLIGENCE", "Public source IP → ASN → organization/network → approximate geolocation.", INDIGO),
        (5, "\U0001F517", "EVIDENCE GRAPH", "Connects Email → Sender → Domain → URL → IP → ASN → Geo → Threat Intel.", NAVY_SOFT),
        (6, "⚖", "EXPLAINABLE RISK", "Multiple evidence signals fused into one transparent, deterministic risk score.", AMBER),
    ]
    cols = 3
    card_w = (content_w - 0.34 * (cols - 1)) / cols
    card_h = 1.55
    gap_y = 0.30
    for idx, (anchor_i, icon, title, desc, col) in enumerate(innov):
        row = idx // cols
        cix = idx % cols
        bx = left_edge + cix * (card_w + 0.34)
        by = board_top + row * (card_h + gap_y)
        ax, ay, ar, _ = centers[anchor_i]
        lx = bx + card_w * (0.2 if cix == 0 else (0.8 if cix == cols - 1 else 0.5))
        add_line_shape(slide, ax, label_bottom, lx, by, color=BORDER_DK, weight=0.85, dash="sysDot")
        add_rect(slide, bx, by, card_w, card_h, fill=CARD, line=BORDER, line_w=0.75, radius=0.09, shadow=True)
        add_rect(slide, bx, by, 0.07, card_h, fill=col, radius=0.5)
        add_oval(slide, bx + 0.22, by + 0.18, 0.42, fill=col)
        add_text(slide, bx + 0.22, by + 0.18, 0.42, 0.42, icon, size=14, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, bx + 0.76, by + 0.20, card_w - 0.9, 0.40, title, size=10, bold=True,
                  color=NAVY_DARK, line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, bx + 0.22, by + 0.70, card_w - 0.42, card_h - 0.82, desc, size=8.6,
                  color=TEXT_MUTED, line_spacing=1.22)

    bar_y = board_top + 2 * (card_h + gap_y) + 0.10
    add_rect(slide, left_edge, bar_y, content_w, 0.42, fill=NAVY_DARK, radius=0.5)
    add_text(slide, left_edge, bar_y, content_w, 0.42,
              "Detection tells you WHAT happened. NetraX investigates WHY — and hands you the evidence.",
              size=10, bold=True, italic=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    footer(slide, 2, question="Why is this more than a spam filter?")
    return slide


# ============================================================== SLIDE 3 ====

def build_slide3(prs):
    slide = add_slide(prs)
    header_band(slide, 3, "Technical Approach", "Email Investigation Pipeline")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    pipe_w = content_w * 0.62
    side_x = left_edge + pipe_w + 0.30
    side_w = content_w - pipe_w - 0.30
    cx0 = left_edge + pipe_w / 2
    top = 1.04

    def stage_label(y, text):
        add_text(slide, left_edge, y, pipe_w, 0.16, text, size=6.8, bold=True, color=TEXT_SOFT, spacing_pt=0.6)

    # 01 ingestion chain
    y1 = top + 0.10
    chip_w, chip_h = 1.7, 0.30
    for i, (icon, lbl) in enumerate([("✉", "SUSPICIOUS EMAIL (.eml / paste / JSON)"), ("\U0001F4E5", "EMAIL PARSER — MIME + indicator extraction")]):
        cy = y1 + i * (chip_h + 0.10)
        add_rect(slide, cx0 - 2.35, cy, 4.7, chip_h, fill=NAVY, radius=0.5)
        add_text(slide, cx0 - 2.2, cy, 0.3, chip_h, icon, size=10, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx0 - 1.85, cy, 4.0, chip_h, lbl, size=7.6, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        if i == 0:
            add_line_shape(slide, cx0, cy + chip_h, cx0, cy + chip_h + 0.10, color=BORDER_DK, weight=1.0)
    stage_label(y1 - 0.16, "01  INGESTION")

    # 02 orchestrator
    row_bottom = y1 + 2 * chip_h + 0.10
    orch_r = 0.40
    orch_cy = row_bottom + 0.30 + orch_r
    add_line_shape(slide, cx0, row_bottom, cx0, orch_cy - orch_r, color=BORDER_DK, weight=1.0)
    add_oval(slide, cx0 - orch_r - 0.08, orch_cy - orch_r - 0.08, (orch_r + 0.08) * 2, fill=NAVY_MED)
    add_oval(slide, cx0 - orch_r, orch_cy - orch_r, orch_r * 2, fill=TEAL_DK)
    add_text(slide, cx0 - orch_r, orch_cy - orch_r + 0.05, orch_r * 2, 0.32, "NETRAX\nAGENT", size=8.2,
              bold=True, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=0.9)
    stage_label(row_bottom + 0.04, "02  ORCHESTRATOR — decides which tools this email needs")

    # 03 tool ecosystem — one row, style marks always-run vs conditional
    tool_y = orch_cy + orch_r + 0.34 + 0.34
    tools = [
        ("\U0001F4C4", "HEADER\nFORENSICS", BLUE, False, ""),
        ("\U0001F517", "URL / DOMAIN\nANALYSIS", BLUE, False, ""),
        ("\U0001F464", "SENDER\nIDENTITY", BLUE, False, ""),
        ("\U0001F6E1", "THREAT INTEL\n(if URL found)", TEAL, True, "PhishTank · URLhaus"),
        ("\U0001F4CD", "GEOLOCATION\n(if public IP)", INDIGO, True, "Source IP → ASN → GeoLite2"),
    ]
    n_t = len(tools)
    t_r = 0.34
    t_gap = 0.42
    t_total = n_t * t_r * 2 + (n_t - 1) * t_gap
    tx0 = cx0 - t_total / 2 + t_r
    max_label_bottom = 0
    for i, (icon, lbl, col, cond, sub) in enumerate(tools):
        tx = tx0 + i * (t_r * 2 + t_gap)
        if cond:
            add_line_shape(slide, cx0, orch_cy + orch_r, tx, tool_y - t_r, color=BORDER_DK, weight=1.0, dash="dash")
            add_oval(slide, tx - t_r, tool_y - t_r, t_r * 2, fill=CARD, line=col, line_w=1.6)
            add_text(slide, tx - t_r, tool_y - t_r, t_r * 2, t_r * 2, icon, size=11, color=col,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            connect_circle_to_circle(slide, cx0, orch_cy, orch_r, tx, tool_y, t_r, color=col, weight=1.3)
            add_oval(slide, tx - t_r, tool_y - t_r, t_r * 2, fill=col)
            add_text(slide, tx - t_r, tool_y - t_r, t_r * 2, t_r * 2, icon, size=11, color=WHITE,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        lbl_color = TEXT_MUTED if cond else NAVY_DARK
        add_text(slide, tx - 0.75, tool_y + t_r + 0.04, 1.5, 0.30, lbl, size=6.8, bold=True,
                  color=lbl_color, italic=cond, align=PP_ALIGN.CENTER, line_spacing=0.9)
        lb = tool_y + t_r + 0.34
        if sub:
            add_text(slide, tx - 0.75, lb + 0.01, 1.5, 0.14, sub, size=5.8, italic=True,
                      color=col, align=PP_ALIGN.CENTER)
            lb += 0.16
        max_label_bottom = max(max_label_bottom, lb)
    stage_label(tool_y - t_r - 0.20, "03  TOOLS")
    legend_x = left_edge
    add_oval(slide, legend_x, tool_y - 0.16, 0.10, fill=BLUE)
    add_text(slide, legend_x + 0.15, tool_y - 0.19, 1.15, 0.16, "always run", size=6.2,
              color=TEXT_MUTED, italic=True)
    add_oval(slide, legend_x, tool_y + 0.06, 0.10, fill=CARD, line=TEAL, line_w=1.2)
    add_text(slide, legend_x + 0.15, tool_y + 0.03, 1.15, 0.16, "conditional", size=6.2,
              color=TEXT_MUTED, italic=True)

    # 04 evidence fusion
    fus_r = 0.26
    fus_cy = max_label_bottom + 0.20 + fus_r
    for i in range(n_t):
        tx = tx0 + i * (t_r * 2 + t_gap)
        add_line_shape(slide, tx, tool_y + t_r, cx0, fus_cy - fus_r, color=BORDER_DK, weight=0.9)
    add_oval(slide, cx0 - fus_r, fus_cy - fus_r, fus_r * 2, fill=NAVY_DARK)
    add_text(slide, cx0 - fus_r, fus_cy - fus_r, fus_r * 2, fus_r * 2, "\U0001F517", size=10, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, cx0 - 1.2, fus_cy + fus_r + 0.03, 2.4, 0.16, "EVIDENCE FUSION", size=7.4, bold=True,
              color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.3)

    # 05 risk engine gauge (4 deterministic levels)
    risk_y = fus_cy + fus_r + 0.16 + 0.16
    add_line_shape(slide, cx0, fus_cy + fus_r + 0.06, cx0, risk_y - 0.02, color=BORDER_DK, weight=1.0)
    gauge_w, gauge_h = pipe_w - 1.9, 0.22
    gx = cx0 - gauge_w / 2
    seg_w = gauge_w / 4
    seg_cols = [GREEN, AMBER, "EA580C", RED]
    seg_lbl = ["LOW", "MEDIUM", "HIGH", "CRITICAL"]
    for i in range(4):
        shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if i in (0, 3) else MSO_SHAPE.RECTANGLE
        add_rect(slide, gx + i * seg_w, risk_y, seg_w, gauge_h, fill=seg_cols[i], radius=0.5, shape_type=shp_type)
        add_text(slide, gx + i * seg_w, risk_y + gauge_h + 0.02, seg_w, 0.14, seg_lbl[i], size=6.2,
                  bold=True, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, left_edge, risk_y - 0.02, gx - left_edge - 0.08, gauge_h,
              "RISK ENGINE\n(0-100, deterministic) →", size=6.8, bold=True, color=NAVY_DARK,
              align=PP_ALIGN.RIGHT, line_spacing=0.95, anchor=MSO_ANCHOR.MIDDLE)

    # 06 findings + report
    dec_y = risk_y + gauge_h + 0.30
    decisions = [("\U0001F4AC", "EXPLAINABLE FINDINGS"), ("\U0001F5C2", "CASE REPORT")]
    n_d = len(decisions)
    d_w = 2.6
    d_gap = 0.4
    dxp = cx0 - (d_w * n_d + d_gap) / 2
    for icon, label in decisions:
        add_rect(slide, dxp, dec_y, d_w, 0.32, fill=BLUE_PALE, radius=0.5, line=BLUE_LIGHT, line_w=0.75)
        add_text(slide, dxp + 0.10, dec_y, 0.4, 0.32, icon, size=10, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, dxp + 0.44, dec_y, d_w - 0.5, 0.32, label, size=8, bold=True, color=BLUE_DK,
                  anchor=MSO_ANCHOR.MIDDLE)
        dxp += d_w + d_gap
    add_text(slide, left_edge, dec_y, 1.6, 0.32, "06  DECIDE", size=6.8, bold=True, color=TEXT_SOFT,
              spacing_pt=0.6, anchor=MSO_ANCHOR.MIDDLE)

    # ---- right sidebar: data + ML ----
    sy = top
    section_label(slide, side_x, sy, "Data Grounding", accent=BLUE, size=10)
    sy += 0.28
    data_items = [
        ("Enron Email Corpus", "structural/language reference — not a labelled class"),
        ("SpamAssassin Corpus", "trains the email content classifier"),
        ("UCI Phishing Websites", "trains the URL classifier"),
        ("PhishTank", "phishing URL threat intel"),
        ("URLhaus", "malware-distribution URL threat intel"),
        ("MaxMind GeoLite2", "IP → ASN → geolocation enrichment"),
    ]
    for name, role in data_items:
        add_multi(slide, side_x, sy, side_w, 0.22, [(name + "  ", NAVY_DARK, True, False), ("— " + role, TEXT_MUTED, False, True)], size=7.6)
        sy += 0.235

    sy += 0.10
    section_label(slide, side_x, sy, "Trained ML Models", accent=TEAL_DK, size=10)
    sy += 0.28
    add_rect(slide, side_x, sy, side_w, 0.62, fill=CARD, line=BORDER, line_w=0.75, radius=0.10)
    add_text(slide, side_x + 0.12, sy + 0.06, side_w - 0.24, 0.18, "Email content classifier", size=8, bold=True, color=NAVY_DARK)
    add_text(slide, side_x + 0.12, sy + 0.25, side_w - 0.24, 0.32,
              "TF-IDF + Linear SVM · F1 = 0.961, ROC-AUC = 0.998\n(measured, held-out SpamAssassin test set)",
              size=6.8, color=TEXT_MUTED, line_spacing=1.1)
    sy += 0.72
    add_rect(slide, side_x, sy, side_w, 0.62, fill=CARD, line=BORDER, line_w=0.75, radius=0.10)
    add_text(slide, side_x + 0.12, sy + 0.06, side_w - 0.24, 0.18, "URL / domain classifier", size=8, bold=True, color=NAVY_DARK)
    add_text(slide, side_x + 0.12, sy + 0.25, side_w - 0.24, 0.32,
              "HistGradientBoosting · F1 = 0.954\n(measured, UCI Phishing Websites test set)",
              size=6.8, color=TEXT_MUTED, line_spacing=1.1)
    sy += 0.80
    section_label(slide, side_x, sy, "Agentic Logic", accent=AMBER, size=10)
    sy += 0.26
    add_bullets(slide, side_x, sy, side_w, 0.9, [
        "No URL in email → URL & threat-intel tools skipped entirely",
        "No public source IP → geolocation skipped entirely",
        "Every skip is logged — auditable, not silent",
    ], size=7.4, line_spacing=1.05, space_after=3, bullet_color=AMBER)
    sy += 0.86
    add_text(slide, side_x, sy, side_w, 0.3,
              "AI orchestrates & explains. The risk engine calculates — never an LLM guessing a number.",
              size=7.2, italic=True, bold=True, color=NAVY_DARK, line_spacing=1.15)

    # ---- bottom tech-stack strip ----
    strip_y = 6.62
    add_line_shape(slide, left_edge, strip_y - 0.06, SLIDE_W_IN - 0.42, strip_y - 0.06, color=BORDER_DK, weight=0.75)
    stacks = [
        ("FRONTEND", "React · TypeScript · Vite · Tailwind CSS · shadcn/ui · Framer Motion · Recharts"),
        ("BACKEND / ML", "Supabase · PostgreSQL · Edge Functions · Python · scikit-learn · Gemini API (explanation layer, planned)"),
    ]
    sw = content_w / 2
    for i, (lbl, val) in enumerate(stacks):
        sx = left_edge + i * sw
        add_text(slide, sx, strip_y, sw - 0.2, 0.16, lbl, size=6.8, bold=True, color=TEXT_SOFT, spacing_pt=0.6)
        add_text(slide, sx, strip_y + 0.17, sw - 0.2, 0.3, val, size=7.6, color=NAVY_DARK, line_spacing=1.05)

    footer(slide, 3, question="No chain-of-thought exposed — only auditable tools, findings and evidence.")
    return slide


# ============================================================== SLIDE 4 ====

def build_slide4(prs):
    slide = add_slide(prs)
    header_band(slide, 4, "Feasibility & Viability", "Prototype → Validation → Deployment-Ready")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.02

    # ---- 3-stage spine ----
    stages = [
        ("PROTOTYPE (TODAY)", NAVY_SOFT, "203 automated tests passing\n(169 TS + 34 Python — measured)"),
        ("VALIDATION", TEAL, "Threshold tuning · cross-validation\nhuman review · feedback loop"),
        ("DEPLOYMENT-READY", BLUE, "Portable Edge Functions · RLS schema\ncredential-gated live intel"),
    ]
    n_s = len(stages)
    arrow_gap = 0.30
    block_w = (content_w - arrow_gap * (n_s - 1)) / n_s
    block_h = 1.10
    bx = left_edge
    for i, (title, col, chips) in enumerate(stages):
        shp = add_rect(slide, bx, top, block_w, block_h, fill=col, radius=0.09)
        add_soft_shadow(shp, blur_pt=8, alpha_pct=15)
        add_text(slide, bx + 0.20, top + 0.14, block_w - 0.4, 0.20, f"0{i+1}", size=10, bold=True, color=WHITE)
        add_text(slide, bx + 0.20, top + 0.36, block_w - 0.4, 0.26, title, size=12.5, bold=True, color=WHITE, spacing_pt=0.3)
        add_text(slide, bx + 0.20, top + 0.66, block_w - 0.4, 0.40, chips, size=7.6, color=WHITE, line_spacing=1.2)
        if i < n_s - 1:
            right_arrow(slide, bx + block_w + 0.02, top + block_h / 2, size=0.20, color=BORDER_DK)
        bx += block_w + arrow_gap

    # ---- challenge -> mitigation rows ----
    map_y = top + block_h + 0.22
    section_label(slide, left_edge, map_y, "Challenge → Mitigation", accent=RED)
    rows_y = map_y + 0.32
    pairs = [
        ("Data quality", "Preprocessing + profiling before training (reports/data_quality_report.md)"),
        ("Threat-intel / API availability", "Adapter architecture; reports “unavailable”, never a fabricated match — cached URLhaus snapshot for demo"),
        ("ML false positives / negatives", "Cross-validated model selection, tuned decision threshold, held-out evaluation"),
        ("Unseen attack patterns", "Multi-signal fusion (headers + URL + ML + threat intel) + a documented feedback loop"),
        ("Computational requirements", "ML inference isolated in its own service; investigation logic is portable/modular"),
        ("Platform security", "Auth + row-level security · input validation · SSRF-safe URL fetches · no secrets in code"),
    ]
    row_h = 0.30
    col1_w = 3.5
    for i, (ch, mit) in enumerate(pairs):
        ry = rows_y + i * (row_h + 0.03)
        add_oval(slide, left_edge, ry + 0.06, 0.14, fill=RED)
        add_text(slide, left_edge + 0.22, ry, col1_w - 0.22, row_h, ch, size=8.6, bold=True, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        right_arrow(slide, left_edge + col1_w + 0.05, ry + row_h / 2, size=0.16, color=BORDER_DK)
        add_oval(slide, left_edge + col1_w + 0.32, ry + 0.06, 0.14, fill=GREEN)
        add_text(slide, left_edge + col1_w + 0.54, ry, content_w - col1_w - 0.6, row_h, mit, size=8.2,
                  color=TEXT_MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.98)

    # ---- why feasible strip ----
    wf_y = rows_y + 6 * (row_h + 0.03) + 0.14
    section_label(slide, left_edge, wf_y, "Why It Is Feasible", accent=GREEN)
    chip_y = wf_y + 0.30
    chips = [
        "Public research datasets already acquired & profiled",
        "PhishTank / URLhaus integrated via tested adapters",
        "GeoLite2 IP→ASN→geo enrichment adapter, unit tested",
        "ML inference kept modular — portable pipeline",
        "Supabase/PostgreSQL schema scales case & evidence storage",
        "Live intel when credentialed; labelled demo data otherwise",
    ]
    cols = 3
    cw = (content_w - 0.24 * (cols - 1)) / cols
    ch_h = 0.46
    for i, txt in enumerate(chips):
        r_, c_ = divmod(i, cols)
        cx = left_edge + c_ * (cw + 0.24)
        cy = chip_y + r_ * (ch_h + 0.10)
        add_rect(slide, cx, cy, cw, ch_h, fill=GREEN_LIGHT, radius=0.5, line="B8E5C4", line_w=0.75)
        add_text(slide, cx + 0.14, cy, 0.26, ch_h, "✓", size=11, bold=True, color=GREEN_DK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx + 0.36, cy, cw - 0.5, ch_h, txt, size=7.6, bold=True, color=GREEN_DK,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)

    footer(slide, 4, question="Built and tested today — not merely proposed.")
    return slide


# ============================================================== SLIDE 5 ====

def build_slide5(prs):
    slide = add_slide(prs)
    header_band(slide, 5, "Impact & Benefits", "Who It Serves & What Changes")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.04

    # ---- ecosystem hub ----
    hub_h = 3.10
    stakeholders = [
        ("Government", "Cyber-crime early-warning signal"),
        ("Banks / Financial Institutions", "Structured phishing/BEC case evidence"),
        ("Enterprises", "Reduced manual triage effort"),
        ("Educational Institutions", "Safer institutional email channels"),
        ("SOC / Security Teams", "Faster triage & evidence correlation"),
        ("Digital Forensic Investigators", "Structured, citable forensic evidence"),
    ]
    outcomes = ["Detection", "Investigation", "Evidence", "Infrastructure Intelligence", "Faster Triage", "Explainable Decisions"]
    n = len(stakeholders)
    row_h = hub_h / n
    lcx = left_edge + 2.55
    rcx = SLIDE_W_IN - 0.42 - 2.05
    hub_cx = SLIDE_W_IN / 2
    hub_cy = top + hub_h / 2
    hub_r = 0.52

    for i, (name, sub) in enumerate(stakeholders):
        cy = top + row_h * (i + 0.5)
        add_line_shape(slide, lcx + 0.10, cy, hub_cx - hub_r, hub_cy, color=BORDER_DK, weight=0.9)
        add_rect(slide, left_edge, cy - 0.20, 2.55, 0.40, fill=CARD, line=BORDER, line_w=0.75, radius=0.5)
        add_text(slide, left_edge + 0.16, cy - 0.20, 2.30, 0.40, name, size=8, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, left_edge, cy + 0.21, 2.55, 0.16, sub, size=6.4, italic=True, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    for i, label in enumerate(outcomes):
        cy = top + row_h * (i + 0.5)
        add_line_shape(slide, hub_cx + hub_r, hub_cy, rcx - 0.10, cy, color=BORDER_DK, weight=0.9)
        add_rect(slide, rcx, cy - 0.16, 2.05, 0.32, fill=BLUE_PALE, line=BLUE_LIGHT, line_w=0.75, radius=0.5)
        add_text(slide, rcx, cy - 0.16, 2.05, 0.32, label, size=7.6, bold=True, color=BLUE_DK,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    add_oval(slide, hub_cx - hub_r - 0.10, hub_cy - hub_r - 0.10, (hub_r + 0.10) * 2, fill=NAVY_MED)
    add_shield(slide, hub_cx - hub_r * 0.72, hub_cy - hub_r * 0.78, hub_r * 1.44, hub_r * 1.44, fill=BLUE, check=True, check_color=WHITE, check_size=hub_r * 1.44)
    add_text(slide, hub_cx - 0.9, hub_cy + hub_r + 0.10, 1.8, 0.2, "NETRAX", size=9.5, bold=True, color=NAVY_DARK, align=PP_ALIGN.CENTER, spacing_pt=0.6)

    add_text(slide, left_edge, top - 0.16, 2.55, 0.16, "WHO IT SERVES", size=7.4, bold=True, color=TEXT_SOFT, spacing_pt=0.8)
    add_text(slide, rcx, top - 0.16, 2.05, 0.16, "WHAT CHANGES", size=7.4, bold=True, color=TEXT_SOFT, align=PP_ALIGN.CENTER, spacing_pt=0.8)

    add_text(slide, left_edge, top + hub_h + 0.02, content_w, 0.22,
              "A decision-support & forensic-intelligence platform — not a replacement for banks, "
              "police or official cybercrime systems.", size=8.4, italic=True, color=TEXT_SOFT,
              align=PP_ALIGN.CENTER)

    # ---- before / after ----
    ba_y = top + hub_h + 0.34
    section_label(slide, left_edge, ba_y, "Before vs. After NetraX", accent=AMBER)
    rows_y = ba_y + 0.32
    before = ["Suspicious\nEmail", "Manual\nInvestigation", "Scattered\nEvidence", "Slow\nDecision"]
    after = ["Suspicious\nEmail", "NetraX\nInvestigation", "Correlated\nEvidence", "Forensic\nRisk", "Actionable\nCase Report"]

    def mini_chain(y, items, col_fill, label, label_col):
        add_text(slide, left_edge, y, 1.3, 0.34, label, size=8, bold=True, color=label_col, anchor=MSO_ANCHOR.MIDDLE)
        n_i = len(items)
        cw = 1.55
        gap = 0.18
        x0 = left_edge + 1.35
        avail = content_w - 1.35
        cw = (avail - gap * (n_i - 1)) / n_i
        for i, it in enumerate(items):
            ix = x0 + i * (cw + gap)
            add_rect(slide, ix, y, cw, 0.34, fill=col_fill, radius=0.5)
            add_text(slide, ix, y, cw, 0.34, it.replace("\n", " "), size=7, bold=True, color=WHITE,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.9)
            if i < n_i - 1:
                right_arrow(slide, ix + cw + 0.01, y + 0.17, size=0.13, color=BORDER_DK)

    mini_chain(rows_y, before, RED, "BEFORE", RED_DK)
    mini_chain(rows_y + 0.46, after, TEAL_DK, "AFTER", TEAL_DK)

    footer(slide, 5, question="This solves a real triage and evidence problem.")
    return slide


# ============================================================== SLIDE 6 ====

def build_slide6(prs):
    slide = add_slide(prs)
    header_band(slide, 6, "Research & References", "Evidence Board")
    left_edge = 0.42
    content_w = SLIDE_W_IN - 0.84
    top = 1.04

    cols = [
        ("DATA & INTELLIGENCE", BLUE, [
            ("CMU Enron Email Dataset", "cs.cmu.edu/~enron"),
            ("Apache SpamAssassin Public Corpus", "spamassassin.apache.org"),
            ("UCI Phishing Websites Dataset", "archive.ics.uci.edu/dataset/327"),
            ("PhishTank", "phishtank.org"),
            ("URLhaus (abuse.ch)", "urlhaus.abuse.ch"),
            ("MaxMind GeoLite2", "dev.maxmind.com/geoip/geolite2-free-geolocation-data"),
        ]),
        ("TECHNOLOGY", TEAL_DK, [
            ("React + TypeScript + Vite", "react.dev"),
            ("Tailwind CSS + shadcn/ui", "tailwindcss.com"),
            ("Supabase / PostgreSQL", "supabase.com/docs"),
            ("Python + scikit-learn", "scikit-learn.org"),
            ("Gemini API", "ai.google.dev/gemini-api/docs"),
            ("(explanation layer — planned, not yet wired)", ""),
        ]),
        ("SECURITY & STANDARDS", INDIGO_DK, [
            ("OWASP Top 10 / SSRF Prevention", "owasp.org"),
            ("NIST Cybersecurity Framework", "nist.gov/cyberframework"),
            ("CERT-In Advisories", "cert-in.org.in"),
            ("Govt. of India Cybercrime Portal", "cybercrime.gov.in"),
        ]),
    ]
    col_w = (content_w - 0.30 * 2) / 3
    col_h = 3.55
    for i, (title, accent, items) in enumerate(cols):
        cx = left_edge + i * (col_w + 0.30)
        add_rect(slide, cx, top, col_w, col_h, fill=CARD, line=BORDER, line_w=0.75, radius=0.09)
        add_rect(slide, cx, top, col_w, 0.06, fill=accent, radius=0.5)
        add_text(slide, cx + 0.20, top + 0.18, col_w - 0.4, 0.22, title, size=10.5, bold=True, color=NAVY_DARK, spacing_pt=0.3)
        iy = top + 0.52
        for name, src in items:
            add_oval(slide, cx + 0.20, iy + 0.05, 0.09, fill=accent)
            add_text(slide, cx + 0.38, iy, col_w - 0.58, 0.20, name, size=8.4, bold=True, color=TEXT_DARK, line_spacing=1.0)
            if src:
                add_text(slide, cx + 0.38, iy + 0.19, col_w - 0.58, 0.16, src, size=7, italic=True, color=TEXT_MUTED)
                iy += 0.44
            else:
                iy += 0.30

    # research basis
    rb_y = top + col_h + 0.22
    section_label(slide, left_edge, rb_y, "Research Basis", accent=AMBER)
    tags = ["Email / Phishing Detection", "URL Phishing Detection", "Email Header Forensics (SPF/DKIM/DMARC)",
            "Threat-Intelligence Correlation", "Agentic / Tool-Use AI Systems", "IP Geolocation Limitations",
            "Explainable Security Systems"]
    tag_y = rb_y + 0.32
    tx = left_edge
    for tag in tags:
        tw = 0.11 * len(tag) + 0.30
        if tx + tw > SLIDE_W_IN - 0.42:
            tx = left_edge
            tag_y += 0.36
        add_rect(slide, tx, tag_y, tw, 0.30, fill=AMBER_LIGHT, radius=0.5, line="F3D9A8", line_w=0.75)
        add_text(slide, tx, tag_y, tw, 0.30, tag, size=7.6, bold=True, color="92400E", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tx += tw + 0.14
    add_text(slide, left_edge, tag_y + 0.40, content_w, 0.2,
              "Topic areas grounding this build — specific citations to be finalized for the final report.",
              size=7.6, italic=True, color=TEXT_SOFT)

    footer(slide, 6, question="Grounded in real, verifiable datasets, tools and standards.")
    return slide


# ============================================================== BUILD ======

def main():
    prs = new_presentation()
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    build_slide4(prs)
    build_slide5(prs)
    build_slide6(prs)
    out = "NetraX_SIH26106_Email_Forensics.pptx"
    prs.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
